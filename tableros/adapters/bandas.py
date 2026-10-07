"""Adapter de la familia `bandas-periodo` (una escala ordenada abierta a lo ancho).

Familia bandas-periodo = una fila por PERIODO y unidad geografica, con una columna por BANDA
de una escala ordenada, y la misma banda repetida en dos unidades (hectareas y porcentaje).
Primera base que la usa: 43 (recursos forrajeros, Observatorio Forrajero Nacional).

Por que es una familia nueva y no una de las nueve de CLAUDE.md:

  - no es `tidy`: ahi 1 fila = 1 observacion y lo que se despivotea son MEDIDAS distintas;
    aca lo que abre a lo ancho es la BANDA, y cada banda viene en dos unidades a la vez;
  - no es `stock`: la 48 tambien abre categorias a lo ancho, pero su grano es el anio y aca
    hay anio + mes + QUINCENA, y ademas la 48 trae una sola unidad (cabezas);
  - no es `wide-mes`: lo que abre a lo ancho no es el mes.

Que hace el adapter, y solo eso:

1. Lee las hojas declaradas en `hojas.datos`, cada una con su fila de header. Las de
   `hojas.ignorar` ni se abren.
2. Renombra columnas origen -> canonico segun `columnas`. Una columna del Excel sin mapeo
   FALLA: nada se descarta en silencio.
3. Despivotea las columnas de `valores` a filas (banda, rol, medida, unidad, valor).
4. Resuelve la geografia contra configs/dims/geo-alias.yaml, igual que el resto de las
   familias. Un departamento sin match FALLA: va al alias, nunca se dropea.
5. Arma la clave de tiempo: anio, mes, quincena, mas un `periodo` legible ("2024-07 1ra") y un
   `orden_periodo` entero para ordenar sin parsear texto.

Lo que NO hace: no valida calidad (eso es qa-datos), no rellena huecos, no corrige un valor
jamas -la columna "Promedio en ha" de la base 43 viene con el mismo numero repetido durante
cuatro anios y medio y entra tal cual, con el caso escrito en el reporte de la base- y no
dropea filas con dato.

Lo unico que descarta son las filas SIN PERIODO Y SIN GEOGRAFIA: al pie de la hoja de la base
43 quedo una fila con dos numeros sueltos y nada mas. No tiene a que observacion pertenecer, y
queda contada en el resumen de la ingesta.

Determinismo: sin timestamps, sin aleatoriedad, y las filas salen en el orden en que estan en
el Excel, hoja por hoja, en el orden del config.
"""
import os
import unicodedata

import openpyxl
import polars as pl
import yaml

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_ALIAS = os.path.join(RAIZ, "configs", "dims", "geo-alias.yaml")

COLUMNAS_SALIDA = [
    "base_id", "entrega", "hoja", "ambito", "fila_origen",
    "grano_tiempo", "anio", "mes", "mes_nombre", "quincena", "periodo", "orden_periodo",
    "provincia", "provincia_nombre", "provincia_id",
    "nivel_geo", "departamento", "geo_nombre", "geo_id", "es_agregado_geo",
    "variable", "banda", "banda_numero", "banda_etiqueta", "banda_rol",
    "medida", "unidad", "agregable", "valor",
    "fuente",
]

ESQUEMA_SALIDA = {
    "base_id": pl.Int64, "entrega": pl.Utf8, "hoja": pl.Utf8, "ambito": pl.Utf8,
    "fila_origen": pl.Int64,
    "grano_tiempo": pl.Utf8, "anio": pl.Int64, "mes": pl.Int64, "mes_nombre": pl.Utf8,
    "quincena": pl.Utf8, "periodo": pl.Utf8, "orden_periodo": pl.Int64,
    "provincia": pl.Utf8, "provincia_nombre": pl.Utf8, "provincia_id": pl.Int64,
    "nivel_geo": pl.Utf8, "departamento": pl.Utf8, "geo_nombre": pl.Utf8, "geo_id": pl.Utf8,
    "es_agregado_geo": pl.Boolean,
    "variable": pl.Utf8, "banda": pl.Utf8, "banda_numero": pl.Int64,
    "banda_etiqueta": pl.Utf8, "banda_rol": pl.Utf8,
    "medida": pl.Utf8, "unidad": pl.Utf8, "agregable": pl.Boolean, "valor": pl.Float64,
    "fuente": pl.Utf8,
}

# Como se numera la quincena dentro del mes. Sale del dato, no de un supuesto: la base 43 usa
# "1ra" y "2da" y nada mas. Una grafia nueva FALLA, para que un cambio de la fuente no entre
# en silencio y rompa el orden de la serie.
QUINCENAS = {"1ra": 1, "2da": 2}


def parse(config, path, entrega=None):
    """Devuelve un polars.DataFrame tidy con las columnas de COLUMNAS_SALIDA."""
    alias = _cargar_alias()
    filas = []
    libro = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for hoja_cfg in config["hojas"]["datos"]:
            filas.extend(_parsear_hoja(config, libro, hoja_cfg, entrega, alias))
    finally:
        libro.close()
    return pl.DataFrame(filas, schema=ESQUEMA_SALIDA).select(COLUMNAS_SALIDA)


def _parsear_hoja(config, libro, hoja_cfg, entrega, alias):
    nombre = hoja_cfg["nombre"]
    if nombre not in libro.sheetnames:
        raise ValueError("La hoja %r no esta en el Excel. Hojas: %s"
                         % (nombre, ", ".join(libro.sheetnames)))
    hoja = libro[nombre]
    header_fila = int(hoja_cfg.get("header_fila", 1))
    posiciones = None
    salida = []
    for nro_fila, celdas in enumerate(hoja.iter_rows(values_only=True), 1):
        if nro_fila < header_fila:
            continue
        if nro_fila == header_fila:
            posiciones = _mapear_header(celdas, config["columnas"], nombre)
            continue
        cruda = {canon: _celda(celdas, pos) for canon, pos in posiciones.items()}
        # Fila sin periodo NI geografia: es el calculo suelto que quedo al pie de la hoja. No
        # pertenece a ninguna observacion, asi que no hay nada que ingerir.
        if cruda.get("anio") is None and _texto(cruda.get("departamento")) is None:
            continue
        salida.extend(_filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega,
                                            nombre, alias))
    if posiciones is None:
        raise ValueError("La hoja %r no tiene la fila de header %d" % (nombre, header_fila))
    return salida


def _mapear_header(celdas, mapeo, nombre_hoja):
    """{canonico: posicion}. Una columna del Excel sin mapeo corta la ingesta."""
    declarado = {str(k).strip(): v for k, v in mapeo.items()}
    posiciones, sin_mapeo = {}, []
    for pos, celda in enumerate(celdas):
        titulo = _texto(celda)
        if titulo is None:
            continue
        if titulo not in declarado:
            sin_mapeo.append(titulo)
            continue
        posiciones[declarado[titulo]] = pos
    if sin_mapeo:
        raise ValueError(
            "La hoja %r trae columnas que el config no mapea: %s. Agregalas a `columnas` "
            "(nada se descarta en silencio)." % (nombre_hoja, ", ".join(sin_mapeo)))
    faltan = sorted(set(declarado.values()) - set(posiciones))
    if faltan:
        raise ValueError("La hoja %r no trae las columnas declaradas: %s"
                         % (nombre_hoja, ", ".join(faltan)))
    return posiciones


def _filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, nombre_hoja, alias):
    geo = _resolver_geo(config, hoja_cfg, cruda, nro_fila, nombre_hoja, alias)
    tiempo = _resolver_tiempo(cruda, nro_fila, nombre_hoja)
    fuente = (config.get("fuente") or {}).get("organismo")
    base_fila = dict(
        base_id=int(config["base"]), entrega=entrega, hoja=nombre_hoja,
        ambito=hoja_cfg.get("ambito"), fila_origen=nro_fila,
        grano_tiempo=(config.get("grano") or {}).get("tiempo", "quincena"),
        fuente=fuente, **tiempo, **geo)
    salida = []
    for canon, declarado in config["valores"].items():
        valor = cruda.get(canon)
        # Un valor que no es numero NO se convierte ni se adivina: entra como nulo y la
        # anomalia queda en el reporte. (En la base 43 pasa en la columna de total, que
        # ademas no se ingiere.)
        if not isinstance(valor, (int, float)) or isinstance(valor, bool):
            valor = None
        salida.append(dict(
            base_fila,
            variable=canon,
            banda=declarado["banda"],
            banda_numero=int(declarado["numero"]),
            banda_etiqueta=declarado["etiqueta"],
            banda_rol=declarado["rol"],
            medida=declarado["medida"],
            unidad=declarado["unidad"],
            agregable=bool(declarado["agregable"]),
            valor=float(valor) if valor is not None else None,
        ))
    return salida


def _resolver_tiempo(cruda, nro_fila, nombre_hoja):
    anio = _entero(cruda.get("anio"))
    mes = _entero(cruda.get("mes"))
    quincena = _texto(cruda.get("quincena"))
    if anio is None or mes is None or quincena is None:
        raise ValueError(
            "Fila %d de la hoja %r: le falta parte del periodo (año=%r, mes=%r, quincena=%r)."
            % (nro_fila, nombre_hoja, anio, mes, quincena))
    numero = QUINCENAS.get(quincena)
    if numero is None:
        raise ValueError(
            "Fila %d de la hoja %r: la quincena %r no es ninguna de las conocidas (%s). "
            "Si la fuente cambio la grafia hay que declararla, no adivinarla."
            % (nro_fila, nombre_hoja, quincena, ", ".join(sorted(QUINCENAS))))
    return {
        "anio": anio,
        "mes": mes,
        "mes_nombre": _texto(cruda.get("mes_nombre")),
        "quincena": quincena,
        # Legible para el sitio y ordenable sin parsear texto. `orden_periodo` es
        # AAAAMMQ: 2024071 es la 1ra quincena de julio de 2024.
        "periodo": "%04d-%02d %s" % (anio, mes, quincena),
        "orden_periodo": anio * 1000 + mes * 10 + numero,
    }


def _cargar_alias():
    with open(RUTA_ALIAS, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _resolver_geo(config, hoja_cfg, cruda, nro_fila, nombre_hoja, alias):
    geo_cfg = dict(config.get("geo") or {})
    geo_cfg.update(hoja_cfg.get("geo") or {})
    normalizar = bool(geo_cfg.get("normalizar_nombre"))

    provincia = _texto(cruda.get("provincia_nombre"))
    departamento = _texto(cruda.get("departamento"))
    if provincia is None:
        raise ValueError("Fila %d de la hoja %r: no trae provincia."
                         % (nro_fila, nombre_hoja))
    provincia_id = (alias.get("provincia") or {}).get(provincia)
    if provincia_id is None:
        raise ValueError(
            "Fila %d de la hoja %r: la provincia %r no esta en geo-alias.yaml. Agregala con "
            "su codigo INDEC (nunca se dropea una fila por geografia)."
            % (nro_fila, nombre_hoja, provincia))

    if departamento is None:
        raise ValueError("Fila %d de la hoja %r: no trae departamento."
                         % (nro_fila, nombre_hoja))

    deptos = (alias.get("departamentos") or {}).get(provincia_id) or {}
    # El nombre CANONICO de cada codigo es el PRIMERO que lo declara en la tabla de alias: ahi
    # la primera entrada es el nombre bueno y las que siguen son sus erratas. Importa porque
    # `geo_nombre` es lo que dim_geo elige como nombre del departamento, colapsando por el
    # minimo alfabetico entre todas las familias. Si esta base aportara su propia grafia
    # ("JUAN F, IBARRA", con coma), ganaria el minimo y la zonificacion de la base 48 -que
    # declara "JUAN F. IBARRA"- dejaria de cerrar. Se aporta el canonico y nada mas.
    canonico = {}
    for nombre_alias, codigo in deptos.items():
        canonico.setdefault(codigo, nombre_alias)
    if normalizar:
        # La tabla de alias esta escrita en mayusculas y sin tildes: se normalizan las dos
        # puntas para que una errata de tipeo no cuente como un departamento nuevo.
        indice = {_clave(k): v for k, v in deptos.items()}
        geo_id = indice.get(_clave(departamento))
    else:
        geo_id = deptos.get(departamento)
    if geo_id is None:
        raise ValueError(
            "Fila %d de la hoja %r: el departamento %r no esta en geo-alias.yaml para la "
            "provincia %s. Agregalo con su codigo INDEC: una fila nunca se dropea por "
            "geografia." % (nro_fila, nombre_hoja, departamento, provincia))
    return {
        "provincia": geo_cfg.get("particion", "sde"),
        "provincia_nombre": provincia,
        "provincia_id": int(provincia_id),
        "nivel_geo": geo_cfg.get("nivel", "departamento"),
        "departamento": departamento,
        "geo_nombre": _normalizar(canonico.get(geo_id, departamento)),
        "geo_id": str(geo_id),
        "es_agregado_geo": False,
    }


def _normalizar(texto):
    """El nombre CANONICO que viaja a dim_geo: mayusculas y sin tildes, con su puntuacion.

    La puntuacion se conserva a proposito. `geo_nombre` es lo que dim_geo elige como nombre del
    departamento, y otras partes del repo lo nombran asi: la zonificacion de la base 48 declara
    "JUAN F. IBARRA" con punto. Sacarle el punto aca hacia que dim_geo se quedara con
    "JUAN F IBARRA" -que ordena antes- y la zonificacion dejaba de cerrar.
    """
    if texto is None:
        return None
    sin_tildes = unicodedata.normalize("NFKD", str(texto))
    sin_tildes = "".join(c for c in sin_tildes if not unicodedata.combining(c))
    return " ".join(sin_tildes.upper().split())


def _clave(texto):
    """La clave con la que se BUSCA en la tabla de alias: ademas, sin puntuacion.

    Es lo que hace que "JUAN F, IBARRA" y "JUAN F. IBARRA" encuentren la misma entrada. Se usa
    solo para buscar; el nombre que se guarda es el de `_normalizar`.
    """
    if texto is None:
        return None
    base = _normalizar(texto)
    limpio = "".join(c if c.isalnum() or c.isspace() else " " for c in base)
    return " ".join(limpio.split())


def _texto(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _celda(celdas, pos):
    return celdas[pos] if pos < len(celdas) else None


def _entero(valor):
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return int(valor)
    texto = str(valor).strip()
    try:
        return int(float(texto))
    except ValueError:
        return None
