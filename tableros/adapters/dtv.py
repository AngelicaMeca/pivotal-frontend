"""Adapter de la familia `dtv` (Declaraciones de Transito Vegetal, SENASA).

Familia dtv = transaccional. Cada fila del Excel es UNA declaracion de transito
vegetal emitida un dia concreto: de que departamento de Santiago sale, a que partido
de que provincia va, que producto, como viene acondicionado y cuanto pesa. El grano
es diario y no hay agregacion previa: los cortes por mes, por anio, por departamento,
por provincia de destino, por tipo de movimiento, por acondicionamiento y por producto
se calculan despues, en marts.

Bases que la usan hoy: 53 (algodon), 56 (batata), 57 (cebolla), 75 (papa). Las cuatro
comparten las 15 columnas de SENASA, pero difieren en cosas que resuelve el CONFIG,
no este archivo: cuantas hojas de datos hay, en que fila esta el header, como se llama
la primera columna ("FECHA EMISION" o "FECHA ") y que trae la segunda ("NRO. DTV",
"ANIO" o "mes").

Que hace el adapter, y solo eso:

1. Lee las hojas declaradas en `hojas.datos`, cada una con su `header_fila`.
2. Renombra columnas origen -> canonico segun `columnas`. Si el Excel trae una columna
   sin mapeo, FALLA: nada se descarta en silencio. Si falta una columna obligatoria
   (`COLUMNAS_OBLIGATORIAS`), tambien falla.
3. Convierte el peso a la unidad canonica del proyecto (toneladas) aplicando, FILA POR
   FILA, el factor que el config declara para la U.M. de esa fila (`unidades.peso.um_a_tn`).
   La U.M. cruda, el peso crudo, la cantidad y el peso unitario viajan igual a la salida
   (`um_origen`, `peso_total_origen`, `cantidad_origen`, `peso_unitario_origen`): la
   conversion nunca pisa el dato de la fuente.
4. Despivotea las medidas declaradas en `valores` a filas (variable, medida, unidad, valor).
5. Resuelve la geografia de las dos puntas contra configs/dims/geo-alias.yaml, igual que
   la familia dte: codigo INDEC solo para las provincias de `geo.resolver` (Santiago), y
   para el resto provincia + partido como texto con geo_id nulo. Un nombre declarado en
   `geo.sin_dato` (el "S/D" de la base 53) queda con geo_id nulo y `geo_sin_dato=true`:
   se agrega a nivel provincia, jamas se dropea la fila.
6. Marca con `es_agregado_fila=true` la fila de totales que algunas hojas traen al pie
   (=SUMA() arrastrado en Excel). Se marca, no se borra.

Lo que NO hace: no valida calidad (eso es qa-datos), no agrega ni netea flujos (eso es
marts), no dropea filas jamas, no corrige el Excel.

Determinismo: sin timestamps, sin aleatoriedad, y las filas salen en el orden en que
estan en el Excel (hoja por hoja, en el orden del config).
"""
import datetime
import os

import openpyxl
import polars as pl
import yaml

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_ALIAS = os.path.join(RAIZ, "configs", "dims", "geo-alias.yaml")

# Columnas canonicas de salida, en orden fijo. El schema es el mismo para toda la familia.
COLUMNAS_SALIDA = [
    "base_id",
    "entrega",
    "hoja",
    "ambito",
    "fila_origen",
    "es_agregado_fila",
    "nro_dtv",
    "grano_tiempo",
    "fecha_hora",
    "fecha",
    "anio",
    "mes",
    "periodo",
    "orden_periodo",
    "provincia",
    "origen_provincia",
    "origen_provincia_id",
    "origen_departamento",
    "origen_geo_id",
    "destino_provincia",
    "destino_provincia_id",
    "destino_departamento",
    "destino_geo_id",
    "alcance",
    "geo_sin_dato",
    "tipo_movimiento",
    "tipo_origen",
    "producto",
    "acondicionamiento",
    "um_origen",
    "um_canonica",
    "factor_tn",
    "cantidad_origen",
    "peso_unitario_origen",
    "peso_total_origen",
    "variable",
    "medida",
    "medida_etiqueta",
    "unidad",
    "agregable",
    "valor",
    "fuente",
]

ESQUEMA_SALIDA = {
    "base_id": pl.Int64,
    "entrega": pl.Utf8,
    "hoja": pl.Utf8,
    "ambito": pl.Utf8,
    "fila_origen": pl.Int64,
    "es_agregado_fila": pl.Boolean,
    "nro_dtv": pl.Utf8,
    "grano_tiempo": pl.Utf8,
    "fecha_hora": pl.Datetime("us"),
    "fecha": pl.Date,
    "anio": pl.Int64,
    "mes": pl.Int64,
    "periodo": pl.Utf8,
    "orden_periodo": pl.Int64,
    "provincia": pl.Utf8,
    "origen_provincia": pl.Utf8,
    "origen_provincia_id": pl.Int64,
    "origen_departamento": pl.Utf8,
    "origen_geo_id": pl.Utf8,
    "destino_provincia": pl.Utf8,
    "destino_provincia_id": pl.Int64,
    "destino_departamento": pl.Utf8,
    "destino_geo_id": pl.Utf8,
    "alcance": pl.Utf8,
    "geo_sin_dato": pl.Boolean,
    "tipo_movimiento": pl.Utf8,
    "tipo_origen": pl.Utf8,
    "producto": pl.Utf8,
    "acondicionamiento": pl.Utf8,
    "um_origen": pl.Utf8,
    "um_canonica": pl.Utf8,
    "factor_tn": pl.Float64,
    "cantidad_origen": pl.Float64,
    "peso_unitario_origen": pl.Float64,
    "peso_total_origen": pl.Float64,
    "variable": pl.Utf8,
    "medida": pl.Utf8,
    "medida_etiqueta": pl.Utf8,
    "unidad": pl.Utf8,
    "agregable": pl.Boolean,
    "valor": pl.Float64,
    "fuente": pl.Utf8,
}

# Sin estas columnas la fila no se puede interpretar. El resto (nro_dtv, mes_declarado,
# anio_declarado, fecha_vencimiento) es opcional y cambia de base en base.
COLUMNAS_OBLIGATORIAS = [
    "fecha_emision", "tipo_movimiento", "origen_provincia", "origen_departamento",
    "destino_provincia", "destino_departamento", "producto", "acondicionamiento",
    "cantidad", "peso_unitario", "um", "peso_total",
]

# Tipos de medida que el config puede pedir en `valores[].origen`.
ORIGENES = ("constante", "columna", "peso_tn")


def parse(config, path, entrega=None):
    """Devuelve un polars.DataFrame tidy con las columnas de COLUMNAS_SALIDA."""
    alias = _cargar_alias()
    filas = []
    libro = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for hoja_cfg in config["hojas"]["datos"]:
            filas.extend(_parsear_hoja(config, libro, hoja_cfg, entrega, alias))
    finally:
        libro.close()
    return pl.DataFrame(filas, schema=ESQUEMA_SALIDA).select(COLUMNAS_SALIDA)


# ---------------------------------------------------------------------------
# Hoja
# ---------------------------------------------------------------------------
def _parsear_hoja(config, libro, hoja_cfg, entrega, alias):
    nombre_hoja = hoja_cfg["nombre"]
    if nombre_hoja not in libro.sheetnames:
        raise ValueError(
            "La hoja %r que declara el config no esta en el archivo. Hojas presentes: %s"
            % (nombre_hoja, libro.sheetnames)
        )
    hoja = libro[nombre_hoja]
    header_fila = int(hoja_cfg.get("header_fila", 1))

    mapeo = dict(config.get("columnas") or {})
    mapeo.update(hoja_cfg.get("columnas") or {})
    mapeo = {_texto(k): v for k, v in mapeo.items()}

    valores = config["valores"]
    fuente = (config.get("indice") or {}).get("fuente")
    peso_cfg = ((config.get("unidades") or {}).get("peso")) or {}
    geo_cfg = config.get("geo") or {}

    filas = []
    posiciones = None
    for nro_fila, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
        if nro_fila < header_fila:
            continue
        if nro_fila == header_fila:
            posiciones = _mapear_header(celdas, mapeo, nombre_hoja)
            continue
        if all(_es_vacia(c) for c in celdas):
            continue  # fila totalmente vacia (relleno del Excel), no es un dato
        cruda = {campo: _celda(celdas, pos) for campo, pos in posiciones.items()}
        filas.extend(
            _filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, valores,
                                  fuente, peso_cfg, geo_cfg, alias)
        )
    return filas


def _mapear_header(celdas, mapeo, nombre_hoja):
    """Header origen -> {campo_canonico: indice_de_columna}. Falla si sobra o falta una."""
    posiciones = {}
    sin_mapeo = []
    for i, celda in enumerate(celdas):
        titulo = _texto(celda)
        if not titulo:
            continue
        if titulo not in mapeo:
            sin_mapeo.append(titulo)
            continue
        posiciones[mapeo[titulo]] = i
    if sin_mapeo:
        raise ValueError(
            "La hoja %r tiene columnas sin mapeo en el config: %s. "
            "Agregalas a `columnas` (o a hojas.datos[].columnas) antes de ingerir."
            % (nombre_hoja, sin_mapeo)
        )
    faltan = [c for c in COLUMNAS_OBLIGATORIAS if c not in posiciones]
    if faltan:
        raise ValueError(
            "La hoja %r no trae las columnas obligatorias de la familia dtv: %s. "
            "Revisa `header_fila` y el mapeo de `columnas` en el config."
            % (nombre_hoja, faltan)
        )
    return posiciones


def _filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, valores,
                          fuente, peso_cfg, geo_cfg, alias):
    """Una fila del Excel (una DTV) -> una fila de salida por cada medida declarada."""
    fecha_hora = _fecha_hora(cruda.get("fecha_emision"), nro_fila, hoja_cfg["nombre"])
    fecha = fecha_hora.date() if fecha_hora else None
    producto = _texto(cruda.get("producto"))

    # Fila de agregado: el =SUMA() que Excel deja al pie de la hoja. No tiene fecha ni
    # producto, solo los totales de alguna columna. Se marca, no se borra.
    es_agregado = fecha is None and producto is None

    geo = _resolver_geo(geo_cfg, cruda, nro_fila, hoja_cfg["nombre"], alias, es_agregado)

    cantidad = _numero(cruda.get("cantidad"), nro_fila, "cantidad", hoja_cfg["nombre"])
    peso_unitario = _numero(cruda.get("peso_unitario"), nro_fila, "peso_unitario",
                            hoja_cfg["nombre"])
    peso_total = _numero(cruda.get("peso_total"), nro_fila, "peso_total", hoja_cfg["nombre"])
    um_origen = _texto(cruda.get("um"))
    um_canonica, factor = _factor_a_tn(um_origen, peso_cfg, nro_fila, hoja_cfg["nombre"],
                                       es_agregado)
    peso_tn = None if (peso_total is None or factor is None) else peso_total * factor

    comun = {
        "base_id": int(config["base"]),
        "entrega": entrega,
        "hoja": hoja_cfg["nombre"],
        "ambito": hoja_cfg.get("ambito"),
        "fila_origen": nro_fila,
        "es_agregado_fila": es_agregado,
        "nro_dtv": _texto(cruda.get("nro_dtv")),
        "grano_tiempo": (config.get("tiempo") or {}).get("grano", "fecha"),
        # SENASA sella la DTV con fecha Y HORA. La hora no se analiza, pero se conserva:
        # sin ella, dos declaraciones distintas del mismo dia, mismo origen, mismo destino
        # y mismo peso quedan como filas identicas y el check de duplicados de qa-datos
        # denuncia 14.223 "repetidas" en la base 53 que en realidad son camiones distintos.
        "fecha_hora": fecha_hora,
        "fecha": fecha,
        "anio": fecha.year if fecha else None,
        "mes": fecha.month if fecha else None,
        "periodo": "%04d-%02d" % (fecha.year, fecha.month) if fecha else None,
        # Mes corrido y CONTIGUO entre diciembre y enero (mismo criterio que adapters/dte.py):
        # con anio*100+mes cualquier check de huecos inventaria 88 meses faltantes.
        "orden_periodo": fecha.year * 12 + (fecha.month - 1) if fecha else None,
        "tipo_movimiento": _texto(cruda.get("tipo_movimiento")),
        "tipo_origen": _texto(cruda.get("tipo_origen")),
        "producto": producto,
        "acondicionamiento": _texto(cruda.get("acondicionamiento")),
        "um_origen": um_origen,
        "um_canonica": um_canonica,
        "factor_tn": factor,
        "cantidad_origen": cantidad,
        "peso_unitario_origen": peso_unitario,
        "peso_total_origen": peso_total,
        "fuente": fuente,
    }
    comun.update(geo)

    disponibles = {"cantidad": cantidad, "peso_unitario": peso_unitario,
                   "peso_total": peso_total}

    salida = []
    for campo, meta in valores.items():
        fila = dict(comun)
        fila.update({
            # `variable` es el nombre canonico de la medida. Junto con (hoja, fila_origen)
            # identifica la celda de origen, asi que es la clave unica de la fila de salida
            # y el criterio de orden de los marts.
            "variable": campo,
            "medida": meta.get("medida"),
            "medida_etiqueta": meta.get("etiqueta"),
            "unidad": meta.get("unidad"),
            "agregable": bool(meta.get("agregable", True)),
            "valor": _valor_de_medida(meta, campo, disponibles, peso_tn, es_agregado,
                                      nro_fila, hoja_cfg["nombre"]),
        })
        salida.append(fila)
    return salida


def _valor_de_medida(meta, campo, disponibles, peso_tn, es_agregado, nro_fila, nombre_hoja):
    """El valor de una medida segun lo que el config declara en `valores[].origen`."""
    origen = meta.get("origen", "columna")
    if origen not in ORIGENES:
        raise ValueError(
            "La medida %r declara origen %r; los validos son %s."
            % (campo, origen, list(ORIGENES))
        )
    if origen == "constante":
        # Contar movimientos: cada fila del Excel es una DTV emitida y vale 1. La fila
        # de total al pie NO es un movimiento, asi que no suma.
        return None if es_agregado else float(meta.get("valor", 1))
    if origen == "peso_tn":
        return peso_tn
    columna = meta.get("columna", campo)
    if columna not in disponibles:
        raise ValueError(
            "Fila %d de la hoja %r: la medida %r pide la columna %r, que la familia dtv "
            "no lee. Columnas disponibles: %s"
            % (nro_fila, nombre_hoja, campo, columna, sorted(disponibles))
        )
    return disponibles[columna]


# ---------------------------------------------------------------------------
# Unidades: PESO TOTAL -> toneladas, segun la U.M. de CADA FILA
# ---------------------------------------------------------------------------
def _factor_a_tn(um_origen, peso_cfg, nro_fila, nombre_hoja, es_agregado=False):
    """(unidad normalizada, factor) para llevar PESO TOTAL de esa fila a toneladas.

    La tabla la declara el config (`unidades.peso.um_a_tn`) porque la misma etiqueta de
    U.M. no significa lo mismo en todas las bases: en la 53 las filas 'U.' vienen con el
    peso total ya en toneladas y en la 56/57/75 vienen en kilos. Eso NO se adivina aca.

    La fila de total al pie no trae U.M. Si en esa base TODAS las U.M. convierten con el
    mismo factor, el =SUMA() del pie esta en esa misma unidad y el config puede decirlo en
    `unidades.peso.factor_fila_total`; asi el total del pie queda comparable contra el
    detalle (es lo que mira el check `fila_total_al_pie` de qa-datos). Si la base mezcla
    factores, el total del pie no esta en ninguna unidad y se deja sin convertir.
    """
    tabla = peso_cfg.get("um_a_tn") or {}
    if um_origen is None:
        if es_agregado and peso_cfg.get("factor_fila_total") is not None:
            return (_texto(peso_cfg.get("canonico")),
                    float(peso_cfg["factor_fila_total"]))
        return None, None            # fila de total al pie: sin U.M., no se convierte
    entrada = tabla.get(um_origen)
    if entrada is None:
        raise ValueError(
            "Fila %d de la hoja %r: la U.M. %r no esta en `unidades.peso.um_a_tn` del "
            "config. Declarala con su factor a toneladas (y avisala en el reporte); "
            "nunca se edita el Excel." % (nro_fila, nombre_hoja, um_origen)
        )
    return _texto(entrada.get("canonica")), float(entrada["factor"])


# ---------------------------------------------------------------------------
# Geografia (dos puntas: origen y destino)
# ---------------------------------------------------------------------------
def _cargar_alias():
    with open(RUTA_ALIAS, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _resolver_geo(geo_cfg, cruda, nro_fila, nombre_hoja, alias, es_agregado):
    resolver = set(geo_cfg.get("resolver") or [])
    propia = geo_cfg.get("provincia_propia")
    sin_dato = set(geo_cfg.get("sin_dato") or [])

    org = _resolver_punta(cruda.get("origen_provincia"), cruda.get("origen_departamento"),
                          resolver, sin_dato, alias, nro_fila, nombre_hoja, "origen",
                          es_agregado)
    dst = _resolver_punta(cruda.get("destino_provincia"), cruda.get("destino_departamento"),
                          resolver, sin_dato, alias, nro_fila, nombre_hoja, "destino",
                          es_agregado)

    if es_agregado:
        alcance = None
    elif org["provincia_id"] == propia and dst["provincia_id"] == propia:
        alcance = "interno"          # el producto se mueve dentro de Santiago del Estero
    elif org["provincia_id"] == propia:
        alcance = "egreso"           # sale de la provincia
    elif dst["provincia_id"] == propia:
        alcance = "ingreso"          # entra a la provincia
    else:
        alcance = "ajeno"            # ninguna punta es la provincia propia

    return {
        "provincia": geo_cfg.get("particion", "sde"),
        "origen_provincia": org["provincia_nombre"],
        "origen_provincia_id": org["provincia_id"],
        "origen_departamento": org["departamento"],
        "origen_geo_id": org["geo_id"],
        "destino_provincia": dst["provincia_nombre"],
        "destino_provincia_id": dst["provincia_id"],
        "destino_departamento": dst["departamento"],
        "destino_geo_id": dst["geo_id"],
        "alcance": alcance,
        "geo_sin_dato": org["sin_dato"] or dst["sin_dato"],
    }


def _resolver_punta(provincia, departamento, resolver, sin_dato, alias, nro_fila,
                    nombre_hoja, punta, es_agregado):
    """Una punta del flujo. Codigo INDEC solo para las provincias que pide el config."""
    provincia = _texto(provincia)
    departamento = _texto(departamento)
    vacia = {"provincia_nombre": None, "provincia_id": None, "departamento": None,
             "geo_id": None, "sin_dato": False}
    if provincia is None:
        if not es_agregado:
            raise ValueError(
                "Fila %d de la hoja %r: no trae provincia de %s y no es la fila de total."
                % (nro_fila, nombre_hoja, punta)
            )
        return vacia

    provincia_id = (alias.get("provincia") or {}).get(provincia)
    if provincia_id is None:
        raise ValueError(
            "Fila %d de la hoja %r: la provincia de %s %r no esta en geo-alias.yaml. "
            "Agregala con su codigo INDEC (nunca se dropea una fila por geografia)."
            % (nro_fila, nombre_hoja, punta, provincia)
        )

    if provincia_id not in resolver:
        # Provincia fuera del alcance del proyecto: se guarda el nombre tal cual viene y
        # el geo_id queda nulo. No se inventa un codigo de partido.
        return {"provincia_nombre": provincia, "provincia_id": provincia_id,
                "departamento": departamento, "geo_id": None, "sin_dato": False}

    if departamento in sin_dato:
        # La fuente dice explicitamente que no sabe el departamento ("S/D"). La fila se
        # conserva entera y se agrega a nivel provincia; inventarle un codigo seria peor.
        return {"provincia_nombre": provincia, "provincia_id": provincia_id,
                "departamento": departamento, "geo_id": None, "sin_dato": True}

    deptos = (alias.get("departamentos") or {}).get(provincia_id) or {}
    if departamento not in deptos:
        raise ValueError(
            "Fila %d de la hoja %r: el departamento de %s %r (provincia %s) no esta en "
            "geo-alias.yaml. Agregalo con su codigo INDEC; jamas se dropea la fila."
            % (nro_fila, nombre_hoja, punta, departamento, provincia)
        )
    return {"provincia_nombre": provincia, "provincia_id": provincia_id,
            "departamento": departamento, "geo_id": str(deptos[departamento]),
            "sin_dato": False}


# ---------------------------------------------------------------------------
# Utilidades chicas
# ---------------------------------------------------------------------------
def _texto(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _es_vacia(celda):
    return celda is None or (isinstance(celda, str) and not celda.strip())


def _celda(celdas, pos):
    return celdas[pos] if pos < len(celdas) else None


def _fecha_hora(valor, nro_fila, nombre_hoja):
    """FECHA EMISION -> datetime. El dia es el grano; la hora viaja para no perder el sello."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime.datetime):
        return valor
    if isinstance(valor, datetime.date):
        return datetime.datetime(valor.year, valor.month, valor.day)
    try:
        return datetime.datetime.fromisoformat(str(valor).strip())
    except ValueError:
        raise ValueError(
            "Fila %d de la hoja %r: la fecha de emision trae %r, que no es una fecha. "
            "Resolvelo en el config, nunca editando el Excel."
            % (nro_fila, nombre_hoja, valor)
        )


def _numero(valor, nro_fila, campo, nombre_hoja):
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        raise ValueError("Fila %d, hoja %r, campo %r: valor booleano"
                         % (nro_fila, nombre_hoja, campo))
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        raise ValueError(
            "Fila %d de la hoja %r: el campo %r trae %r, que no es numero. "
            "Resolvelo en el config, nunca editando el Excel."
            % (nro_fila, nombre_hoja, campo, valor)
        )
