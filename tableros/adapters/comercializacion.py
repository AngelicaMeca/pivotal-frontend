"""Adapter de la familia `comercializacion` (Siocarnes: hacienda vendida para faena).

Una fila del Excel = una LINEA DE COMERCIALIZACION: tal dia, de tal departamento de origen,
hacia tal zona de destino, de tal raza y tal categoria, se vendieron N cabezas que pesaron K
kilos y se pagaron a tanto el kilo.

Por que es una familia y no una de las que ya estan:
  - No es `dte` ni `dtv`: esas son documentos de transito, su unidad es el movimiento y su
    geografia es origen-destino entre provincias. Aca no hay documento ni provincia de
    destino, hay una ZONA de destino de SENASA ("ZONA 10") y, sobre todo, hay dos cosas que
    ninguna de las dos familias tiene: la RAZA y el PRECIO.
  - No es `stock`: ahi la fila es una existencia a una fecha, no una transaccion.
  - El precio es lo que obliga a separar. Es la primera familia del repo que trae un valor NO
    AGREGABLE junto a los agregables en la misma fila: sumar precios no significa nada y el
    promedio hay que ponderarlo por kilos. Se marca con `agregable: false` y los marts tienen
    que respetarlo.

La salida es tidy-larga, igual que el resto: una fila por cada columna de valor declarada en
el config, con la raza y la categoria como dimensiones.
"""
import os
import unicodedata

import openpyxl
import polars as pl
import yaml

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_ALIAS = os.path.join(RAIZ, "configs", "dims", "geo-alias.yaml")

COLUMNAS_SALIDA = [
    "base_id",
    "entrega",
    "hoja",
    "ambito",
    "fila_origen",
    "grano_tiempo",
    "anio",
    "mes",
    "provincia",
    "provincia_nombre",
    "provincia_id",
    "nivel_geo",
    "departamento",
    "geo_nombre",
    "geo_id",
    "es_agregado_geo",
    "zona_destino",
    "raza",
    "especie",
    "categoria",
    "variable",
    "medida",
    "medida_etiqueta",
    "unidad",
    "agregable",
    "valor",
    "fuente",
]

ESQUEMA_SALIDA = {
    "base_id": pl.Int64,
    "entrega": pl.Utf8,
    "hoja": pl.Utf8,
    "ambito": pl.Utf8,
    "fila_origen": pl.Int64,
    "grano_tiempo": pl.Utf8,
    "anio": pl.Int64,
    "mes": pl.Int64,
    "provincia": pl.Utf8,
    "provincia_nombre": pl.Utf8,
    "provincia_id": pl.Int64,
    "nivel_geo": pl.Utf8,
    "departamento": pl.Utf8,
    "geo_nombre": pl.Utf8,
    "geo_id": pl.Utf8,
    "es_agregado_geo": pl.Boolean,
    "zona_destino": pl.Utf8,
    "raza": pl.Utf8,
    "especie": pl.Utf8,
    "categoria": pl.Utf8,
    "variable": pl.Utf8,
    "medida": pl.Utf8,
    "medida_etiqueta": pl.Utf8,
    "unidad": pl.Utf8,
    "agregable": pl.Boolean,
    "valor": pl.Float64,
    "fuente": pl.Utf8,
}


def parse(config, path, entrega=None):
    alias = _cargar_alias()
    libro = openpyxl.load_workbook(path, read_only=True, data_only=True)
    filas = []
    try:
        for hoja_cfg in (config.get("hojas") or {}).get("datos") or []:
            filas.extend(_parsear_hoja(config, libro, hoja_cfg, entrega, alias))
    finally:
        libro.close()
    if not filas:
        return pl.DataFrame(schema=ESQUEMA_SALIDA)
    return pl.DataFrame(filas, schema=ESQUEMA_SALIDA).select(COLUMNAS_SALIDA)


def _parsear_hoja(config, libro, hoja_cfg, entrega, alias):
    nombre_hoja = hoja_cfg["nombre"]
    if nombre_hoja not in libro.sheetnames:
        raise ValueError(
            "La hoja %r que declara el config no esta en el archivo. Hojas presentes: %s"
            % (nombre_hoja, libro.sheetnames)
        )
    hoja = libro[nombre_hoja]
    header_fila = int(hoja_cfg.get("header_fila", 1))

    mapeo = dict(config.get("columnas") or {})
    mapeo.update(hoja_cfg.get("columnas") or {})
    mapeo = {_texto(k): v for k, v in mapeo.items()}

    valores = config["valores"]
    fuente = (config.get("indice") or {}).get("fuente")
    especie = config.get("especie")
    factor_tn = float((config.get("unidades") or {}).get("kg_a_tn", 1000.0))
    um_esperadas = set((config.get("unidades") or {}).get("esperadas") or [])

    filas = []
    posiciones = None
    for nro_fila, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
        if nro_fila < header_fila:
            continue
        if nro_fila == header_fila:
            posiciones = _mapear_header(celdas, mapeo, nombre_hoja)
            continue
        if all(_es_vacia(c) for c in celdas):
            continue
        cruda = {campo: _celda(celdas, pos) for campo, pos in posiciones.items()}
        # El encabezado REPETIDO adentro de los datos: las hojas de Siocarnes lo traen cuando
        # la exportacion se hizo por tramos. No es una observacion y no se cuenta como fila
        # mala; se saltea y se sigue.
        if _texto(cruda.get("fecha")) == _texto(_clave_header(mapeo, "fecha")):
            continue
        if _es_vacia(cruda.get("origen_departamento")):
            continue  # pie de tabla
        filas.extend(
            _filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, valores,
                                  fuente, alias, especie, factor_tn, um_esperadas)
        )
    return filas


def _clave_header(mapeo, campo):
    """El texto del Excel que mapea a `campo`. Sirve para reconocer el header repetido."""
    for origen, destino in mapeo.items():
        if destino == campo:
            return origen
    return None


def _mapear_header(celdas, mapeo, nombre_hoja):
    posiciones = {}
    sin_mapeo = []
    for i, celda in enumerate(celdas):
        nombre = _texto(celda)
        if nombre is None:
            continue
        if nombre not in mapeo:
            sin_mapeo.append(nombre)
            continue
        destino = mapeo[nombre]
        if destino is not None:
            posiciones[destino] = i
    if sin_mapeo:
        raise ValueError(
            "La hoja %r tiene columnas sin mapeo en el config: %s. Agregalas a `columnas` "
            "antes de ingerir." % (nombre_hoja, sin_mapeo)
        )
    return posiciones


def _filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, valores, fuente,
                          alias, especie, factor_tn, um_esperadas):
    anio = _entero(cruda.get("anio"))
    if anio is None:
        anio = _anio_de_fecha(cruda.get("fecha"))
    if anio is None:
        anio = _entero(hoja_cfg.get("anio_fijo"))
    if anio is None:
        raise ValueError(
            "Fila %d de la hoja %r: no se pudo determinar el anio."
            % (nro_fila, hoja_cfg["nombre"])
        )
    mes = _entero(cruda.get("mes"))

    um = _texto(cruda.get("um"))
    if um_esperadas and um is not None and um not in um_esperadas:
        raise ValueError(
            "Fila %d de la hoja %r: la unidad de medida %r no estaba prevista. Las "
            "declaradas son %s; una unidad nueva cambia como se convierten los kilos."
            % (nro_fila, hoja_cfg["nombre"], um, sorted(um_esperadas))
        )

    geo = _resolver_geo(config, cruda, nro_fila, hoja_cfg["nombre"], alias)
    comun = {
        "base_id": int(config["base"]),
        "entrega": entrega,
        "hoja": hoja_cfg["nombre"],
        "ambito": hoja_cfg.get("ambito"),
        "fila_origen": nro_fila,
        "grano_tiempo": "anio_mes" if mes is not None else "anio",
        "anio": anio,
        "mes": mes,
        "zona_destino": _texto(cruda.get("zona_destino")),
        "raza": _texto(cruda.get("raza")),
        "especie": especie,
        "categoria": _texto(cruda.get("categoria")),
        "fuente": fuente,
    }
    comun.update(geo)

    salida = []
    for campo, meta in valores.items():
        origen = meta.get("origen", "columna")
        if origen == "constante":
            valor = float(meta.get("valor", 1))
        else:
            columna = meta.get("columna", campo)
            if columna not in cruda:
                continue
            valor = _numero(cruda.get(columna), nro_fila, columna, hoja_cfg["nombre"])
            if valor is not None and origen == "kg_a_tn":
                valor = valor / factor_tn
        fila = dict(comun)
        fila.update({
            "variable": campo,
            "medida": meta.get("medida"),
            "medida_etiqueta": meta.get("etiqueta"),
            "unidad": meta.get("unidad"),
            # Un PRECIO no se suma. Se marca y los marts tienen que respetarlo: el promedio
            # de un precio va ponderado por los kilos, nunca es el promedio de la columna.
            "agregable": bool(meta.get("agregable", True)),
            "valor": valor,
        })
        salida.append(fila)
    return salida


# ---------------------------------------------------------------------------
# Geografia (la misma punta que las otras familias)
# ---------------------------------------------------------------------------
def _cargar_alias():
    with open(RUTA_ALIAS, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _resolver_geo(config, cruda, nro_fila, nombre_hoja, alias):
    geo_cfg = config.get("geo") or {}
    provincia = _texto(geo_cfg.get("provincia_fija"))
    departamento = _texto(cruda.get("origen_departamento"))
    provincia_id = (alias.get("provincia") or {}).get(provincia)
    if provincia_id is None:
        raise ValueError(
            "La provincia %r de `geo.provincia_fija` no esta en geo-alias.yaml." % provincia
        )

    deptos = (alias.get("departamentos") or {}).get(provincia_id) or {}
    # El nombre CANONICO de cada codigo (la primera entrada que lo declara en geo-alias). Es
    # lo que se publica: una base nunca le impone su grafia al resto del repo.
    canonico = {}
    for nombre_alias, codigo in deptos.items():
        canonico.setdefault(str(codigo), _normalizar(nombre_alias))
    indice = {_normalizar(k): v for k, v in deptos.items()}
    clave = _normalizar(departamento)
    if clave not in indice:
        raise ValueError(
            "Fila %d de la hoja %r: el departamento de origen %r no esta en geo-alias.yaml. "
            "Agregalo con su codigo INDEC; jamas se dropea la fila."
            % (nro_fila, nombre_hoja, departamento)
        )
    geo_id = str(indice[clave])
    return {
        "provincia": geo_cfg.get("particion", "sde"),
        "provincia_nombre": provincia,
        "provincia_id": provincia_id,
        "nivel_geo": "departamento",
        "departamento": departamento,
        "geo_nombre": canonico.get(geo_id, _normalizar(departamento)),
        "geo_id": geo_id,
        "es_agregado_geo": False,
    }


# ---------------------------------------------------------------------------
# Utilidades chicas
# ---------------------------------------------------------------------------
def _normalizar(texto):
    if texto is None:
        return None
    sin_tildes = "".join(
        c for c in unicodedata.normalize("NFD", str(texto))
        if unicodedata.category(c) != "Mn"
    )
    return " ".join(sin_tildes.upper().split())


def _texto(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _es_vacia(celda):
    return celda is None or (isinstance(celda, str) and not celda.strip())


def _celda(celdas, pos):
    return celdas[pos] if pos < len(celdas) else None


def _entero(valor):
    if _es_vacia(valor):
        return None
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return None


def _anio_de_fecha(valor):
    if valor is None:
        return None
    if hasattr(valor, "year"):
        return int(valor.year)
    texto = str(valor).strip()
    if len(texto) >= 4 and texto[:4].isdigit():
        return int(texto[:4])
    return None


def _numero(valor, nro_fila, campo, nombre_hoja):
    """Numero o nulo. Un texto en una columna de valor FALLA: no se convierte a cero."""
    if _es_vacia(valor):
        return None
    if isinstance(valor, bool):
        raise ValueError(
            "Fila %d de la hoja %r: la columna %r trae un booleano."
            % (nro_fila, nombre_hoja, campo)
        )
    try:
        return float(valor)
    except (TypeError, ValueError):
        raise ValueError(
            "Fila %d de la hoja %r: la columna %r trae %r, que no es un numero. "
            "No se convierte a cero: el dato se revisa en el Excel."
            % (nro_fila, nombre_hoja, campo, valor)
        )
