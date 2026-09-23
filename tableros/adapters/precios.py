"""Adapter de la familia `precios` (tablas planas de precios por producto).

Familia precios = una fila del Excel es UNA observacion de precio: un dia, un producto
descrito por varias dimensiones (grupo, especie, variedad, envase, calidad, tamanio) y
un precio unitario. No hay geografia por fila mas alla del ORIGEN de la mercaderia (la
columna Procedencia), y no hay volumen comercializado.

Bases que la usan hoy: 8 (Precios MCBA origen Santiago, fuente MAGyP sobre el Mercado
Central de Buenos Aires).

LO QUE HACE DISTINTA A ESTA FAMILIA (leer antes de tocar nada): su medida NO ES
ADITIVA. Todas las medidas que tenia el pipeline hasta hoy (toneladas, cabezas,
hectareas, movimientos) se suman; un precio no se suma, se promedia. Por eso:

  * cada medida declara en el config su `agregacion` ('promedio' | 'suma') y el adapter
    FALLA si una medida no aditiva no la declara. La regla viaja en cada fila de
    staging (columna `agregacion`), asi que ningun mart ni ninguna vista tiene que
    adivinarla;
  * cada medida declara tambien su `ponderacion`. En la base 8 es la frase "sin ponderar:
    la fuente no informa volumen comercializado", y esa frase viaja pegada al dato hasta
    el pie del grafico. El promedio mensual es el promedio simple de los promedios
    diarios: NO es el precio medio de lo que se vendio en el mes. Es una limitacion real
    del dato, no una decision del pipeline, y cualquier vista que lo publique tiene que
    poder decirlo.

Que hace el adapter, y solo eso:

1. Lee las hojas declaradas en `hojas.datos`, cada una con su `header_fila`.
2. Renombra columnas origen -> canonico segun `columnas`. Si el Excel trae una columna
   sin mapeo, FALLA: nada se descarta en silencio. Si falta una columna obligatoria
   (`COLUMNAS_OBLIGATORIAS`), tambien falla.
3. Normaliza las dimensiones de producto segun `dimensiones` (mayusculas, mapa de
   equivalencias y clasificacion). De cada dimension salen TRES columnas, siempre las
   mismas y sin excepciones por base: `<dim>` (valor canonico), `<dim>_origen` (lo que
   dice el Excel, intacto) y `<dim>_clase` (la clase que le asigna el config, o nulo).
   Asi una inconsistencia de la fuente -la base 8 escribe "BANDEJA" 689 veces y
   "Bandeja" 309- se unifica para poder agrupar, sin perder lo que estaba escrito.
4. Resuelve la procedencia contra configs/dims/geo-alias.yaml a nivel PROVINCIA, con el
   diccionario `geo.procedencias` del config. Nunca se dropea una fila por geografia.
5. Despivotea las medidas declaradas en `valores` a filas (variable, medida, unidad,
   agregacion, ponderacion, valor).
6. Conserva el anio y el mes que el Excel trae escritos aparte de la fecha
   (`anio_declarado`, `mes_declarado`) y marca con `tiempo_declarado_coincide=false` las
   filas donde no coinciden con la fecha. La fuente de verdad es la FECHA; la
   discrepancia se cuenta, no se corrige.

Lo que NO hace: no valida calidad (eso es qa-datos), no promedia ni agrega nada (eso es
marts y las vistas), no dropea filas jamas, no corrige el Excel.

Determinismo: sin timestamps, sin aleatoriedad, y las filas salen en el orden en que
estan en el Excel (hoja por hoja, en el orden del config).
"""
import datetime
import os

import openpyxl
import polars as pl
import yaml

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_ALIAS = os.path.join(RAIZ, "configs", "dims", "geo-alias.yaml")

# Las dimensiones de producto de la familia, en orden fijo. Una base que no tenga alguna
# simplemente no la mapea en `columnas` y la columna queda nula: no se agregan campos
# nuevos por base (eso es lo que hace que la familia sea una familia).
DIMENSIONES = ("grupo", "especie", "variedad", "envase", "calidad", "tamanio")

# Columnas canonicas de salida, en orden fijo.
COLUMNAS_SALIDA = [
    "base_id",
    "entrega",
    "hoja",
    "ambito",
    "fila_origen",
    "es_agregado_fila",
    "grano_tiempo",
    "fecha",
    "anio",
    "mes",
    "dia",
    "periodo",
    "orden_periodo",
    "anio_declarado",
    "mes_declarado",
    "tiempo_declarado_coincide",
    "provincia",
    "mercado",
    "mercado_nombre",
    "procedencia_origen",
    "origen_provincia",
    "origen_provincia_id",
    "origen_geo_id",
    "nivel_geo",
    "geo_sin_dato",
]
for _dim in DIMENSIONES:
    COLUMNAS_SALIDA += [_dim, "%s_origen" % _dim, "%s_clase" % _dim]
COLUMNAS_SALIDA += [
    "variable",
    "medida",
    "medida_etiqueta",
    "unidad",
    "moneda",
    "expresion_precio",
    "agregable",
    "agregacion",
    "ponderacion",
    "valor",
    "fuente",
]

ESQUEMA_SALIDA = {
    "base_id": pl.Int64,
    "entrega": pl.Utf8,
    "hoja": pl.Utf8,
    "ambito": pl.Utf8,
    "fila_origen": pl.Int64,
    "es_agregado_fila": pl.Boolean,
    "grano_tiempo": pl.Utf8,
    "fecha": pl.Date,
    "anio": pl.Int64,
    "mes": pl.Int64,
    "dia": pl.Int64,
    "periodo": pl.Utf8,
    "orden_periodo": pl.Int64,
    "anio_declarado": pl.Int64,
    "mes_declarado": pl.Utf8,
    "tiempo_declarado_coincide": pl.Boolean,
    "provincia": pl.Utf8,
    "mercado": pl.Utf8,
    "mercado_nombre": pl.Utf8,
    "procedencia_origen": pl.Utf8,
    "origen_provincia": pl.Utf8,
    "origen_provincia_id": pl.Int64,
    "origen_geo_id": pl.Utf8,
    "nivel_geo": pl.Utf8,
    "geo_sin_dato": pl.Boolean,
    "variable": pl.Utf8,
    "medida": pl.Utf8,
    "medida_etiqueta": pl.Utf8,
    "unidad": pl.Utf8,
    "moneda": pl.Utf8,
    "expresion_precio": pl.Utf8,
    "agregable": pl.Boolean,
    "agregacion": pl.Utf8,
    "ponderacion": pl.Utf8,
    "valor": pl.Float64,
    "fuente": pl.Utf8,
}
for _dim in DIMENSIONES:
    ESQUEMA_SALIDA[_dim] = pl.Utf8
    ESQUEMA_SALIDA["%s_origen" % _dim] = pl.Utf8
    ESQUEMA_SALIDA["%s_clase" % _dim] = pl.Utf8

# Sin estas dos columnas la fila no se puede interpretar: sin fecha no hay serie y sin
# especie no hay producto. El resto de las dimensiones es opcional y cambia de base en base.
COLUMNAS_OBLIGATORIAS = ["fecha", "especie"]

# Tipos de medida que el config puede pedir en `valores[].origen`.
ORIGENES = ("columna", "constante")

# Formas de normalizar una dimension (`dimensiones.<dim>.normalizar`).
NORMALIZACIONES = ("ninguna", "mayusculas", "minusculas", "capitalizar")

# Reglas de agregacion que el config puede declarar en `valores[].agregacion`. El
# pipeline no sabe hacer otras: si aparece una base que pide mediana o promedio
# ponderado, se agrega aca y en los marts a la vez, nunca solo en un YAML.
AGREGACIONES = ("suma", "promedio")


def parse(config, path, entrega=None):
    """Devuelve un polars.DataFrame tidy con las columnas de COLUMNAS_SALIDA."""
    alias = _cargar_alias()
    filas = []
    libro = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for hoja_cfg in config["hojas"]["datos"]:
            filas.extend(_parsear_hoja(config, libro, hoja_cfg, entrega, alias))
    finally:
        libro.close()
    return pl.DataFrame(filas, schema=ESQUEMA_SALIDA).select(COLUMNAS_SALIDA)


# ---------------------------------------------------------------------------
# Hoja
# ---------------------------------------------------------------------------
def _parsear_hoja(config, libro, hoja_cfg, entrega, alias):
    nombre_hoja = hoja_cfg["nombre"]
    if nombre_hoja not in libro.sheetnames:
        raise ValueError(
            "La hoja %r que declara el config no esta en el archivo. Hojas presentes: %s"
            % (nombre_hoja, libro.sheetnames)
        )
    hoja = libro[nombre_hoja]
    header_fila = int(hoja_cfg.get("header_fila", 1))

    mapeo = dict(config.get("columnas") or {})
    mapeo.update(hoja_cfg.get("columnas") or {})
    mapeo = {_texto(k): v for k, v in mapeo.items()}

    valores = _valores_validados(config)
    dims_cfg = config.get("dimensiones") or {}
    geo_cfg = config.get("geo") or {}
    mercado_cfg = config.get("mercado") or {}
    unidad_cfg = ((config.get("unidades") or {}).get("precio")) or {}
    fuente = (config.get("indice") or {}).get("fuente")

    filas = []
    posiciones = None
    for nro_fila, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
        if nro_fila < header_fila:
            continue
        if nro_fila == header_fila:
            posiciones = _mapear_header(celdas, mapeo, nombre_hoja)
            continue
        if all(_es_vacia(c) for c in celdas):
            continue  # fila totalmente vacia (relleno del Excel), no es un dato
        cruda = {campo: _celda(celdas, pos) for campo, pos in posiciones.items()}
        filas.extend(
            _filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, valores,
                                  dims_cfg, geo_cfg, mercado_cfg, unidad_cfg, fuente, alias)
        )
    return filas


def _mapear_header(celdas, mapeo, nombre_hoja):
    """Header origen -> {campo_canonico: indice_de_columna}. Falla si sobra o falta una."""
    posiciones = {}
    sin_mapeo = []
    for i, celda in enumerate(celdas):
        titulo = _texto(celda)
        if not titulo:
            continue
        if titulo not in mapeo:
            sin_mapeo.append(titulo)
            continue
        posiciones[mapeo[titulo]] = i
    if sin_mapeo:
        raise ValueError(
            "La hoja %r tiene columnas sin mapeo en el config: %s. "
            "Agregalas a `columnas` (o a hojas.datos[].columnas) antes de ingerir."
            % (nombre_hoja, sin_mapeo)
        )
    faltan = [c for c in COLUMNAS_OBLIGATORIAS if c not in posiciones]
    if faltan:
        raise ValueError(
            "La hoja %r no trae las columnas obligatorias de la familia precios: %s. "
            "Revisa `header_fila` y el mapeo de `columnas` en el config."
            % (nombre_hoja, faltan)
        )
    return posiciones


def _valores_validados(config):
    """`valores` del config, con la regla de agregacion chequeada una vez por corrida."""
    valores = config["valores"]
    for campo, meta in valores.items():
        agregacion = meta.get("agregacion")
        if agregacion is None:
            raise ValueError(
                "La medida %r no declara `agregacion`. En la familia precios es "
                "obligatoria: un precio se promedia y un conteo se suma, y el pipeline "
                "no lo adivina. Poné agregacion: promedio o agregacion: suma en el config."
                % campo
            )
        if agregacion not in AGREGACIONES:
            raise ValueError(
                "La medida %r declara agregacion %r; las que el pipeline sabe hacer son %s."
                % (campo, agregacion, list(AGREGACIONES))
            )
        if agregacion == "suma" and not meta.get("agregable", True):
            raise ValueError(
                "La medida %r dice agregable: false y agregacion: suma a la vez. "
                "Si no es aditiva no se suma." % campo
            )
        if agregacion == "promedio" and meta.get("agregable", True):
            raise ValueError(
                "La medida %r se agrega por promedio, asi que tiene que declarar "
                "agregable: false (no se suma)." % campo
            )
    return valores


def _filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, valores, dims_cfg,
                          geo_cfg, mercado_cfg, unidad_cfg, fuente, alias):
    """Una fila del Excel (una observacion de precio) -> una fila por cada medida."""
    fecha = _fecha(cruda.get("fecha"), nro_fila, hoja_cfg["nombre"])
    dims = {d: _normalizar_dim(d, cruda.get(d), dims_cfg.get(d) or {}) for d in DIMENSIONES}

    # Fila de agregado: si alguna hoja trajera un =SUMA() al pie no tendria ni fecha ni
    # producto. Se marca, no se borra. En la base 8 no hay ninguna.
    es_agregado = fecha is None and all(v["valor"] is None for v in dims.values())

    geo = _resolver_geo(geo_cfg, cruda.get("procedencia"), nro_fila, hoja_cfg["nombre"],
                        alias, es_agregado)

    anio_declarado = _entero(cruda.get("anio_declarado"), nro_fila, "anio_declarado",
                             hoja_cfg["nombre"])
    mes_declarado = _texto(cruda.get("mes_declarado"))
    coincide = _tiempo_coincide(fecha, anio_declarado, mes_declarado, config)

    comun = {
        "base_id": int(config["base"]),
        "entrega": entrega,
        "hoja": hoja_cfg["nombre"],
        "ambito": hoja_cfg.get("ambito"),
        "fila_origen": nro_fila,
        "es_agregado_fila": es_agregado,
        "grano_tiempo": (config.get("tiempo") or {}).get("grano", "fecha"),
        "fecha": fecha,
        "anio": fecha.year if fecha else None,
        "mes": fecha.month if fecha else None,
        "dia": fecha.day if fecha else None,
        "periodo": "%04d-%02d" % (fecha.year, fecha.month) if fecha else None,
        # Mes corrido y CONTIGUO entre diciembre y enero (mismo criterio que adapters/dte.py
        # y adapters/dtv.py): con anio*100+mes cualquier check de huecos inventa meses.
        "orden_periodo": fecha.year * 12 + (fecha.month - 1) if fecha else None,
        "anio_declarado": anio_declarado,
        "mes_declarado": mes_declarado,
        "tiempo_declarado_coincide": coincide,
        "mercado": mercado_cfg.get("codigo"),
        "mercado_nombre": mercado_cfg.get("nombre"),
        "moneda": unidad_cfg.get("moneda"),
        "expresion_precio": unidad_cfg.get("expresion"),
        "fuente": fuente,
    }
    comun.update(geo)
    for dim, res in dims.items():
        comun[dim] = res["valor"]
        comun["%s_origen" % dim] = res["origen"]
        comun["%s_clase" % dim] = res["clase"]

    disponibles = {
        "precio": _numero(cruda.get("precio"), nro_fila, "precio", hoja_cfg["nombre"]),
    }

    salida = []
    for campo, meta in valores.items():
        fila = dict(comun)
        fila.update({
            # `variable` es el nombre canonico de la medida. Junto con (hoja, fila_origen)
            # identifica la celda de origen, asi que es la clave unica de la fila de salida
            # y el criterio de orden de los marts.
            "variable": campo,
            "medida": meta.get("medida"),
            "medida_etiqueta": meta.get("etiqueta"),
            "unidad": meta.get("unidad"),
            "agregable": bool(meta.get("agregable", True)),
            "agregacion": meta.get("agregacion"),
            # Como se pondera esa agregacion, en castellano llano. Viaja pegado al dato
            # para que el pie del grafico lo pueda decir sin volver al config.
            "ponderacion": meta.get("ponderacion"),
            "valor": _valor_de_medida(meta, campo, disponibles, es_agregado, nro_fila,
                                      hoja_cfg["nombre"]),
        })
        salida.append(fila)
    return salida


def _valor_de_medida(meta, campo, disponibles, es_agregado, nro_fila, nombre_hoja):
    """El valor de una medida segun lo que el config declara en `valores[].origen`."""
    origen = meta.get("origen", "columna")
    if origen not in ORIGENES:
        raise ValueError(
            "La medida %r declara origen %r; los validos son %s."
            % (campo, origen, list(ORIGENES))
        )
    if origen == "constante":
        # Contar observaciones: cada fila del Excel es una cotizacion y vale 1. Una fila
        # de agregado no es una observacion, asi que no cuenta.
        return None if es_agregado else float(meta.get("valor", 1))
    columna = meta.get("columna", campo)
    if columna not in disponibles:
        raise ValueError(
            "Fila %d de la hoja %r: la medida %r pide la columna %r, que la familia "
            "precios no lee. Columnas disponibles: %s"
            % (nro_fila, nombre_hoja, campo, columna, sorted(disponibles))
        )
    return disponibles[columna]


# ---------------------------------------------------------------------------
# Dimensiones de producto
# ---------------------------------------------------------------------------
def _normalizar_dim(dim, valor, cfg):
    """{valor canonico, valor del Excel, clase}. Todo lo decide el config, nada el codigo."""
    origen = _texto(valor)
    if origen is None:
        return {"valor": None, "origen": None, "clase": None}

    forma = cfg.get("normalizar", "ninguna")
    if forma not in NORMALIZACIONES:
        raise ValueError(
            "La dimension %r declara normalizar: %r; las formas validas son %s."
            % (dim, forma, list(NORMALIZACIONES))
        )
    if forma == "mayusculas":
        canonico = origen.upper()
    elif forma == "minusculas":
        canonico = origen.lower()
    elif forma == "capitalizar":
        canonico = origen.capitalize()
    else:
        canonico = origen

    # `mapa` es para equivalencias de verdad (dos nombres distintos para la misma cosa),
    # no para erratas de mayusculas: eso ya lo resolvio `normalizar`.
    canonico = (cfg.get("mapa") or {}).get(canonico, canonico)
    clases = cfg.get("clases") or {}
    clase = clases.get(canonico, cfg.get("clase_por_defecto"))
    return {"valor": canonico, "origen": origen, "clase": clase}


# ---------------------------------------------------------------------------
# Geografia: la procedencia de la mercaderia, a nivel provincia
# ---------------------------------------------------------------------------
def _cargar_alias():
    with open(RUTA_ALIAS, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def _resolver_geo(geo_cfg, procedencia, nro_fila, nombre_hoja, alias, es_agregado):
    """Procedencia del Excel -> provincia INDEC. Jamas se dropea una fila por geografia."""
    procedencia = _texto(procedencia)
    base = {
        "provincia": geo_cfg.get("particion", "sde"),
        "nivel_geo": geo_cfg.get("nivel", "provincia"),
        "procedencia_origen": procedencia,
        "origen_provincia": None,
        "origen_provincia_id": None,
        "origen_geo_id": None,
        "geo_sin_dato": True,
    }
    if procedencia is None:
        if not es_agregado:
            raise ValueError(
                "Fila %d de la hoja %r: no trae procedencia y no es una fila de total."
                % (nro_fila, nombre_hoja)
            )
        return base

    if procedencia in (geo_cfg.get("sin_dato") or []):
        # La fuente dice explicitamente que no sabe el origen. La fila se conserva entera.
        return base

    procedencias = geo_cfg.get("procedencias") or {}
    if procedencia not in procedencias:
        raise ValueError(
            "Fila %d de la hoja %r: la procedencia %r no esta declarada en "
            "`geo.procedencias` del config. Agregala apuntando al nombre de provincia "
            "tal como figura en configs/dims/geo-alias.yaml (nunca se dropea una fila "
            "por geografia)." % (nro_fila, nombre_hoja, procedencia)
        )
    nombre = procedencias[procedencia]
    provincia_id = (alias.get("provincia") or {}).get(nombre)
    if provincia_id is None:
        raise ValueError(
            "Fila %d de la hoja %r: la provincia %r (procedencia %r) no esta en "
            "geo-alias.yaml. Agregala con su codigo INDEC."
            % (nro_fila, nombre_hoja, nombre, procedencia)
        )
    base.update({
        "origen_provincia": nombre,
        "origen_provincia_id": int(provincia_id),
        "origen_geo_id": str(provincia_id),
        "geo_sin_dato": False,
    })
    return base


# ---------------------------------------------------------------------------
# Tiempo
# ---------------------------------------------------------------------------
def _tiempo_coincide(fecha, anio_declarado, mes_declarado, config):
    """El Excel repite anio y mes aparte de la fecha: se compara, no se corrige.

    Devuelve None cuando no hay con que comparar. La abreviatura de mes de cada base la
    declara el config (`tiempo.meses_declarados`), porque es texto de la fuente.
    """
    if fecha is None:
        return None
    if anio_declarado is None and mes_declarado is None:
        return None
    meses = (config.get("tiempo") or {}).get("meses_declarados") or {}
    ok = True
    if anio_declarado is not None:
        ok = ok and anio_declarado == fecha.year
    if mes_declarado is not None and meses:
        ok = ok and meses.get(mes_declarado.strip().lower()) == fecha.month
    return ok


# ---------------------------------------------------------------------------
# Utilidades chicas
# ---------------------------------------------------------------------------
def _texto(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _es_vacia(celda):
    return celda is None or (isinstance(celda, str) and not celda.strip())


def _celda(celdas, pos):
    return celdas[pos] if pos < len(celdas) else None


def _fecha(valor, nro_fila, nombre_hoja):
    """La columna Fecha -> date. El grano de la familia es el dia."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    try:
        return datetime.date.fromisoformat(str(valor).strip())
    except ValueError:
        raise ValueError(
            "Fila %d de la hoja %r: la fecha trae %r, que no es una fecha. "
            "Resolvelo en el config, nunca editando el Excel."
            % (nro_fila, nombre_hoja, valor)
        )


def _numero(valor, nro_fila, campo, nombre_hoja):
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        raise ValueError("Fila %d, hoja %r, campo %r: valor booleano"
                         % (nro_fila, nombre_hoja, campo))
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("$", "").replace(" ", "")
    texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        raise ValueError(
            "Fila %d de la hoja %r: el campo %r trae %r, que no es numero. "
            "Resolvelo en el config, nunca editando el Excel."
            % (nro_fila, nombre_hoja, campo, valor)
        )


def _entero(valor, nro_fila, campo, nombre_hoja):
    numero = _numero(valor, nro_fila, campo, nombre_hoja)
    return None if numero is None else int(numero)
