"""Adapter de la familia `tidy`.

Familia tidy = el Excel ya viene largo y limpio: header en una fila, 1 fila = 1
observacion, y (en el mejor caso) ids INDEC de provincia y departamento en el propio
archivo. Es el formato objetivo del pipeline. Primera base que lo usa: 9 (cultivos
extensivos, MAGyP).

Que hace el adapter, y solo eso:

1. Lee las hojas declaradas en el config (`hojas.datos`), con la fila de header que
   diga el config. Las hojas de `hojas.ignorar` ni se abren.
2. Renombra columnas origen -> canonico segun `columnas` (el mapeo de la hoja se
   mergea sobre el global). Si aparece una columna del Excel sin mapeo, FALLA: nada
   se descarta en silencio.
3. Despivotea las columnas de `medidas` a filas (medida, valor, unidad).
4. Clasifica la geografia (departamento / total provincial / pais) y el cultivo
   (agregado / componente / simple). Marca, nunca borra ni corrige.

Lo que NO hace: no valida calidad (eso es qa-datos), no calcula rankings ni
variaciones (eso es marts), no dropea filas jamas.

Determinismo: sin timestamps, sin aleatoriedad, y las filas salen en el orden en que
estan en el Excel (hoja por hoja, en el orden del config).
"""
import openpyxl
import polars as pl

# Columnas canonicas de salida, en orden fijo. El schema es el mismo para toda la familia.
COLUMNAS_SALIDA = [
    "base_id",
    "entrega",
    "hoja",
    "ambito",
    "fila_origen",
    "provincia",
    "provincia_id",
    "provincia_nombre",
    "nivel_geo",
    "geo_id",
    "geo_nombre",
    "departamento_id",
    "departamento",
    "es_agregado_geo",
    "grano_tiempo",
    "campania",
    "campania_inicio",
    "campania_fin",
    "cultivo",
    "cultivo_grupo",
    "cultivo_rol",
    "medida",
    "medida_etiqueta",
    "unidad",
    "agregable",
    "valor",
    "fuente",
]

ESQUEMA_SALIDA = {
    "base_id": pl.Int64,
    "entrega": pl.Utf8,
    "hoja": pl.Utf8,
    "ambito": pl.Utf8,
    "fila_origen": pl.Int64,
    "provincia": pl.Utf8,
    "provincia_id": pl.Int64,
    "provincia_nombre": pl.Utf8,
    "nivel_geo": pl.Utf8,
    "geo_id": pl.Utf8,
    "geo_nombre": pl.Utf8,
    "departamento_id": pl.Int64,
    "departamento": pl.Utf8,
    "es_agregado_geo": pl.Boolean,
    "grano_tiempo": pl.Utf8,
    "campania": pl.Utf8,
    "campania_inicio": pl.Int64,
    "campania_fin": pl.Int64,
    "cultivo": pl.Utf8,
    "cultivo_grupo": pl.Utf8,
    "cultivo_rol": pl.Utf8,
    "medida": pl.Utf8,
    "medida_etiqueta": pl.Utf8,
    "unidad": pl.Utf8,
    "agregable": pl.Boolean,
    "valor": pl.Float64,
    "fuente": pl.Utf8,
}


def parse(config, path, entrega=None):
    """Devuelve un polars.DataFrame tidy con las columnas de COLUMNAS_SALIDA."""
    filas = []
    libro = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for hoja_cfg in config["hojas"]["datos"]:
            filas.extend(_parsear_hoja(config, libro, hoja_cfg, entrega))
    finally:
        libro.close()
    return pl.DataFrame(filas, schema=ESQUEMA_SALIDA).select(COLUMNAS_SALIDA)


# ---------------------------------------------------------------------------
# Hoja
# ---------------------------------------------------------------------------
def _parsear_hoja(config, libro, hoja_cfg, entrega):
    nombre_hoja = hoja_cfg["nombre"]
    if nombre_hoja not in libro.sheetnames:
        raise ValueError(
            "La hoja %r que declara el config no esta en el archivo. Hojas presentes: %s"
            % (nombre_hoja, libro.sheetnames)
        )
    hoja = libro[nombre_hoja]
    header_fila = int(hoja_cfg.get("header_fila", 1))

    mapeo = dict(config.get("columnas") or {})
    mapeo.update(hoja_cfg.get("columnas") or {})
    mapeo = {_texto(k): v for k, v in mapeo.items()}

    medidas = config["medidas"]
    fuente = (config.get("indice") or {}).get("fuente")
    grano = (config.get("tiempo") or {}).get("grano", "campania")

    filas = []
    posiciones = None
    for nro_fila, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
        if nro_fila < header_fila:
            continue
        if nro_fila == header_fila:
            posiciones = _mapear_header(celdas, mapeo, nombre_hoja)
            continue
        if all(_es_vacia(c) for c in celdas):
            continue  # fila totalmente vacia (relleno del Excel), no es un dato
        cruda = {campo: _celda(celdas, pos) for campo, pos in posiciones.items()}
        filas.extend(_filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, medidas, fuente, grano))
    return filas


def _mapear_header(celdas, mapeo, nombre_hoja):
    """Header origen -> {campo_canonico: indice_de_columna}. Falla si sobra una columna."""
    posiciones = {}
    sin_mapeo = []
    for i, celda in enumerate(celdas):
        titulo = _texto(celda)
        if not titulo:
            continue
        if titulo not in mapeo:
            sin_mapeo.append(titulo)
            continue
        posiciones[mapeo[titulo]] = i
    if sin_mapeo:
        raise ValueError(
            "La hoja %r tiene columnas sin mapeo en el config: %s. "
            "Agregalas a `columnas` (o a hojas.datos[].columnas) antes de ingerir."
            % (nombre_hoja, sin_mapeo)
        )
    return posiciones


def _filas_de_observacion(config, hoja_cfg, cruda, nro_fila, entrega, medidas, fuente, grano):
    """Una fila del Excel -> una fila de salida por cada medida declarada."""
    geo = _resolver_geo(config, hoja_cfg, cruda, nro_fila)
    campania = _texto(cruda.get("campania"))
    inicio, fin = _partir_campania(campania, nro_fila, hoja_cfg["nombre"])
    cultivo = _texto(cruda.get("cultivo"))
    grupo, rol = _clasificar_cultivo(cultivo, config.get("agregados_cultivo"))

    salida = []
    for medida, meta in medidas.items():
        if medida not in cruda:
            continue  # la hoja no trae esa medida; se anota por ausencia, no se inventa
        fila = {
            "base_id": int(config["base"]),
            "entrega": entrega,
            "hoja": hoja_cfg["nombre"],
            "ambito": hoja_cfg.get("ambito"),
            "fila_origen": nro_fila,
            "grano_tiempo": grano,
            "campania": campania,
            "campania_inicio": inicio,
            "campania_fin": fin,
            "cultivo": cultivo,
            "cultivo_grupo": grupo,
            "cultivo_rol": rol,
            "medida": medida,
            "medida_etiqueta": meta.get("etiqueta"),
            "unidad": meta.get("unidad"),
            "agregable": bool(meta.get("agregable", True)),
            "valor": _numero(cruda.get(medida), nro_fila, medida, hoja_cfg["nombre"]),
            "fuente": fuente,
        }
        fila.update(geo)
        salida.append(fila)
    return salida


# ---------------------------------------------------------------------------
# Geografia
# ---------------------------------------------------------------------------
def _resolver_geo(config, hoja_cfg, cruda, nro_fila):
    geo_cfg = config.get("geo") or {}
    geo_hoja = hoja_cfg.get("geo") or {}
    nivel_hoja = geo_hoja.get("nivel", geo_cfg.get("nivel", "departamento"))

    if nivel_hoja == "pais":
        return {
            "provincia": geo_hoja.get("provincia_slug", "ar"),
            "provincia_id": None,
            "provincia_nombre": None,
            "nivel_geo": "pais",
            "geo_id": geo_hoja.get("geo_id_fijo", "AR"),
            "geo_nombre": _texto(cruda.get("geo_nombre")),
            "departamento_id": None,
            "departamento": None,
            "es_agregado_geo": True,
        }

    provincia_id = _entero(cruda.get("provincia_id"))
    provincias = geo_cfg.get("provincias") or {}
    if provincia_id not in provincias:
        raise ValueError(
            "Fila %d de la hoja %r: idProvincia %r no esta en geo.provincias del config. "
            "Agregalo (nunca se dropea una fila por geografia)."
            % (nro_fila, hoja_cfg["nombre"], provincia_id)
        )
    provincia_nombre = _texto(cruda.get("provincia_nombre"))
    departamento_id = _entero(cruda.get("departamento_id"))
    departamento = _texto(cruda.get("departamento"))

    agregados = geo_cfg.get("agregados") or {}
    es_total = departamento_id == agregados.get("departamento_id_total") or departamento in (
        agregados.get("nombres_total") or []
    )
    if es_total:
        # Total provincial embebido entre el detalle. Se marca y se guarda la etiqueta
        # original en geo_nombre; departamento queda en null para que no se cuele en un
        # group by departamental.
        return {
            "provincia": provincias[provincia_id],
            "provincia_id": provincia_id,
            "provincia_nombre": provincia_nombre,
            "nivel_geo": "provincia",
            "geo_id": str(provincia_id),
            "geo_nombre": departamento or provincia_nombre,
            "departamento_id": None,
            "departamento": None,
            "es_agregado_geo": True,
        }

    return {
        "provincia": provincias[provincia_id],
        "provincia_id": provincia_id,
        "provincia_nombre": provincia_nombre,
        "nivel_geo": "departamento",
        "geo_id": "%d%03d" % (provincia_id, departamento_id),
        "geo_nombre": departamento,
        "departamento_id": departamento_id,
        "departamento": departamento,
        "es_agregado_geo": False,
    }


# ---------------------------------------------------------------------------
# Tiempo (grano campania)
# ---------------------------------------------------------------------------
def _partir_campania(campania, nro_fila, nombre_hoja):
    """"2014/15" -> (2014, 2015). "2000/01" -> (2000, 2001)."""
    if not campania:
        return None, None
    partes = campania.split("/")
    if len(partes) != 2 or not partes[0].isdigit() or not partes[1].isdigit():
        raise ValueError(
            "Fila %d de la hoja %r: campania %r no tiene el formato AAAA/AA."
            % (nro_fila, nombre_hoja, campania)
        )
    inicio = int(partes[0])
    fin = inicio + 1
    if str(fin)[-2:] != partes[1].zfill(2):
        raise ValueError(
            "Fila %d de la hoja %r: campania %r no son dos anios seguidos."
            % (nro_fila, nombre_hoja, campania)
        )
    return inicio, fin


# ---------------------------------------------------------------------------
# Cultivos: agregado vs componente
# ---------------------------------------------------------------------------
def _clasificar_cultivo(cultivo, cfg):
    """(grupo, rol) con rol en {agregado, componente, simple}.

    Regla general: el nombre que termina en el sufijo declarado (" total") es el
    agregado del grupo; los nombres listados como componentes son sus partes. Sirve
    para soja 1ra/2da/total, trigo, poroto y cebada sin escribir un caso por cultivo.
    """
    if not cultivo:
        return None, None
    if not cfg:
        return cultivo, "simple"
    grupos = cfg.get("grupos") or {}
    for grupo, detalle in grupos.items():
        if cultivo == detalle.get("agregado"):
            return grupo, "agregado"
    for grupo, detalle in grupos.items():
        if cultivo in (detalle.get("componentes") or []):
            return grupo, "componente"
    sufijo = cfg.get("sufijo_agregado") or ""
    if sufijo and cultivo.lower().endswith(sufijo.lower()):
        return cultivo[: -len(sufijo)].strip(), "agregado"
    return cultivo, "simple"


# ---------------------------------------------------------------------------
# Utilidades chicas
# ---------------------------------------------------------------------------
def _texto(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _es_vacia(celda):
    return celda is None or (isinstance(celda, str) and not celda.strip())


def _celda(celdas, pos):
    return celdas[pos] if pos < len(celdas) else None


def _entero(valor):
    if valor is None or valor == "":
        return None
    return int(valor)


def _numero(valor, nro_fila, medida, nombre_hoja):
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        raise ValueError("Fila %d, hoja %r, medida %r: valor booleano" % (nro_fila, nombre_hoja, medida))
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        raise ValueError(
            "Fila %d de la hoja %r: la medida %r trae %r, que no es numero. "
            "Resolvelo en el config, nunca editando el Excel."
            % (nro_fila, nombre_hoja, medida, valor)
        )
