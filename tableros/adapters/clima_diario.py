"""Adapter de la familia `clima-diario` (estaciones meteorologicas: una hoja por estacion,
una fila por DIA, una columna por variable).

Primera base: la 111 (estaciones agrometeorologicas del INTA en Santiago del Estero).

Forma del Excel que espera (todo lo especifico lo declara el config de la base):

  * una hoja de ESTACIONES (`hojas.estaciones`): una fila por estacion debajo del encabezado,
    con su Id, nombre, coordenadas, etc. Termina en la primera fila sin Id; lo que haya debajo
    (en la 111, las definiciones de variables que escribio JC) no se lee.
  * una hoja de DATOS por estacion (`hojas.datos`): una columna de fecha y una columna por
    variable. Cada hoja declara a que estacion pertenece con `estacion_id` (el Id de la hoja de
    estaciones): nunca se deduce del nombre de la hoja.

Salida: formato LARGO, una fila por estacion x fecha x variable, con nulos EXPLICITOS (cada
variable que trae la hoja tiene su fila todos los dias, haya dato o no). Nulo es nulo: una
celda vacia nunca se convierte a cero. No se redondea nada: el redondeo es de presentacion.

Que corta la ingesta (para que nada cambie en silencio):
  * una hoja del libro que el config no declara (ni como datos, ni como estaciones, ni en
    `ignorar`) o una hoja declarada que no esta;
  * un encabezado de columna que no figura en el diccionario de variables
    (`variables: configs/dims/...yaml`, campo `columnas_origen`): columna nueva = drift;
  * una fecha que no se entiende, o una fecha REPETIDA en la hoja que el config no declara en
    `fechas_repetidas` (y, si la declara, que la fila descartada contradiga a la conservada);
  * un valor numerico que no es numero, o un texto de una variable de categoria fuera de su
    `dominio`;
  * una estacion sin hoja, una hoja sin estacion, o un departamento que no esta en
    configs/dims/geo-alias.yaml.

Las variables de tipo `categoria` (la direccion del viento: 'NE', 'C ') van a `valor_texto`
sin los espacios de relleno; las numericas van a `valor`.
"""
import datetime as dt
import os
import re
import unicodedata

import openpyxl
import polars as pl
import yaml

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_ALIAS = os.path.join(RAIZ, "configs", "dims", "geo-alias.yaml")

TIPOS_DE_VALOR = ("numero", "categoria")

# "2010-12-14 00:00:00.0" (como lo exporta el INTA) o "2010-12-14" a secas. El grano es el
# DIA: una hora distinta de la medianoche corta la ingesta, salvo las que el config declare en
# `tiempo.horas_aceptadas` con su motivo (en la 111, las 01:00 de los dias en que empezaba el
# horario de verano y la medianoche no existio).
RE_FECHA = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?: (\d{2}:\d{2}:\d{2})(?:\.0+)?)?$")
HORA_DEL_DIA = "00:00:00"
RE_FECHA_DMA = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")

ESQUEMA_SALIDA = {
    "base_id": pl.Int64, "entrega": pl.Utf8, "hoja": pl.Utf8, "ambito": pl.Utf8,
    "fila_origen": pl.Int64, "grano_tiempo": pl.Utf8, "fecha": pl.Date,
    "anio": pl.Int64, "mes": pl.Int64, "dia": pl.Int64,
    "provincia": pl.Utf8, "provincia_id": pl.Int64,
    "estacion_id": pl.Int64, "estacion": pl.Utf8, "estacion_nombre": pl.Utf8,
    "estacion_tipo": pl.Utf8, "estacion_id_interno": pl.Utf8,
    "localidad": pl.Utf8, "ubicacion": pl.Utf8,
    "latitud": pl.Float64, "longitud": pl.Float64, "altura": pl.Float64,
    "estacion_desde_declarado": pl.Date, "estacion_hasta_declarado": pl.Date,
    "nivel_geo": pl.Utf8, "geo_id": pl.Utf8, "geo_nombre": pl.Utf8,
    "geo_verificar": pl.Boolean,
    "variable": pl.Utf8, "variable_origen": pl.Utf8, "variable_etiqueta": pl.Utf8,
    "unidad": pl.Utf8, "tipo_valor": pl.Utf8,
    "valor": pl.Float64, "valor_texto": pl.Utf8, "fuente": pl.Utf8,
}
COLUMNAS_SALIDA = list(ESQUEMA_SALIDA)


def parse(config, path, entrega=None):
    variables = _cargar_variables(config)
    alias = _cargar_alias()
    columnas = {c: [] for c in COLUMNAS_SALIDA}
    libro = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        _verificar_hojas(config, libro)
        estaciones = _leer_estaciones(config, libro)
        _verificar_estaciones(config, estaciones)
        for hoja_cfg in config["hojas"]["datos"]:
            estacion = estaciones[int(hoja_cfg["estacion_id"])]
            geo = _resolver_geo(config, hoja_cfg, alias)
            _parsear_hoja(config, libro, hoja_cfg, estacion, geo, variables, entrega,
                          columnas)
    finally:
        libro.close()
    return pl.DataFrame(columnas, schema=ESQUEMA_SALIDA).select(COLUMNAS_SALIDA)


# ---------------------------------------------------------------------------
# Hojas y estaciones
# ---------------------------------------------------------------------------
def _verificar_hojas(config, libro):
    hojas = config.get("hojas") or {}
    declaradas = [h["nombre"] for h in hojas.get("datos") or []]
    declaradas.append((hojas.get("estaciones") or {}).get("nombre"))
    ignoradas = [h["nombre"] if isinstance(h, dict) else h for h in hojas.get("ignorar") or []]
    faltan = [h for h in declaradas if h not in libro.sheetnames]
    if faltan:
        raise ValueError(
            "Base %s: el config declara hojas que no estan en el archivo: %s. Hojas presentes: "
            "%s" % (config["base"], faltan, libro.sheetnames)
        )
    sobran = [h for h in libro.sheetnames if h not in declaradas and h not in ignoradas]
    if sobran:
        raise ValueError(
            "Base %s: el archivo trae hojas que el config no declara: %s. Si es una estacion "
            "nueva, agregala a `hojas.datos`; si no es dato, a `hojas.ignorar`."
            % (config["base"], sobran)
        )


def _leer_estaciones(config, libro):
    """{estacion_id: {campo: valor}} de la hoja de estaciones."""
    cfg = config["hojas"]["estaciones"]
    hoja = libro[cfg["nombre"]]
    header_fila = int(cfg.get("header_fila", 1))
    mapeo = {_texto(k): v for k, v in (cfg.get("columnas") or {}).items()}
    estaciones, posiciones = {}, None
    for nro_fila, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
        if nro_fila < header_fila:
            continue
        if nro_fila == header_fila:
            posiciones = {}
            for i, celda in enumerate(celdas):
                nombre = _texto(celda)
                if nombre is None:
                    continue
                if nombre not in mapeo:
                    raise ValueError(
                        "Hoja %r: la columna %r no esta en `hojas.estaciones.columnas` del "
                        "config." % (cfg["nombre"], nombre)
                    )
                posiciones[mapeo[nombre]] = i
            if "estacion_id" not in posiciones:
                raise ValueError("Hoja %r: falta la columna del Id de la estacion."
                                 % cfg["nombre"])
            continue
        id_celda = _celda(celdas, posiciones["estacion_id"])
        if _es_vacia(id_celda):
            break  # fin de la tabla: lo que sigue son notas de JC, no estaciones
        registro = {campo: _celda(celdas, pos) for campo, pos in posiciones.items()}
        estacion_id = _entero_estricto(id_celda, nro_fila, "Id", cfg["nombre"])
        if estacion_id in estaciones:
            raise ValueError("Hoja %r: el Id de estacion %s aparece dos veces."
                             % (cfg["nombre"], estacion_id))
        estaciones[estacion_id] = {
            "estacion_nombre": _texto(registro.get("estacion_nombre")),
            "estacion_tipo": _texto(registro.get("estacion_tipo")),
            "estacion_id_interno": _texto(registro.get("estacion_id_interno")),
            "localidad": _texto(registro.get("localidad")),
            "ubicacion": _texto(registro.get("ubicacion")),
            "latitud": _numero(registro.get("latitud"), nro_fila, "Latitud", cfg["nombre"]),
            "longitud": _numero(registro.get("longitud"), nro_fila, "Longitud", cfg["nombre"]),
            "altura": _numero(registro.get("altura"), nro_fila, "Altura", cfg["nombre"]),
            "desde": _fecha_dma(registro.get("desde_declarado"), nro_fila, cfg["nombre"]),
            "hasta": _fecha_dma(registro.get("hasta_declarado"), nro_fila, cfg["nombre"]),
        }
    return estaciones


def _verificar_estaciones(config, estaciones):
    por_hoja = [int(h["estacion_id"]) for h in config["hojas"]["datos"]]
    sin_estacion = [i for i in por_hoja if i not in estaciones]
    if sin_estacion:
        raise ValueError(
            "Base %s: hojas de datos con `estacion_id` que no esta en la hoja de estaciones: %s"
            % (config["base"], sin_estacion)
        )
    sin_hoja = [i for i in sorted(estaciones) if i not in por_hoja]
    if sin_hoja:
        raise ValueError(
            "Base %s: estaciones de la hoja de estaciones sin hoja de datos declarada: %s"
            % (config["base"], sin_hoja)
        )
    repetidas = sorted({i for i in por_hoja if por_hoja.count(i) > 1})
    if repetidas:
        raise ValueError("Base %s: dos hojas declaran la misma estacion: %s"
                         % (config["base"], repetidas))


def _resolver_geo(config, hoja_cfg, alias):
    """(provincia_id, geo_id, geo_nombre canonico, verificar) del departamento de la hoja."""
    geo_cfg = config.get("geo") or {}
    nombre_prov = _texto(geo_cfg.get("provincia_fija"))
    prov_id = (alias.get("provincia") or {}).get(nombre_prov)
    if prov_id is None:
        raise ValueError("La provincia %r de `geo.provincia_fija` no esta en geo-alias.yaml."
                         % nombre_prov)
    departamentos = (alias.get("departamentos") or {}).get(prov_id) or {}
    # Nombre canonico de cada codigo = la PRIMERA entrada que lo declara (regla de dim_geo:
    # una base nueva nunca le impone su propia grafia al resto del sitio).
    canonico = {}
    for nombre, codigo in departamentos.items():
        canonico.setdefault(codigo, nombre)
    indice = {_normalizar(n): c for n, c in departamentos.items()}
    depto = hoja_cfg.get("departamento")
    codigo = indice.get(_normalizar(depto)) if depto else None
    if codigo is None:
        raise ValueError(
            "Hoja %r: el departamento %r no esta en configs/dims/geo-alias.yaml. Agregalo con "
            "su codigo INDEC y `verificar: true`; nunca se descartan filas por geografia."
            % (hoja_cfg["nombre"], depto)
        )
    return prov_id, str(codigo), canonico[codigo], bool(hoja_cfg.get("verificar", False))


# ---------------------------------------------------------------------------
# Hoja de datos diarios
# ---------------------------------------------------------------------------
def _parsear_hoja(config, libro, hoja_cfg, estacion, geo, variables, entrega, columnas):
    nombre_hoja = hoja_cfg["nombre"]
    hoja = libro[nombre_hoja]
    header_fila = int(hoja_cfg.get("header_fila", 1))
    tiempo_cfg = config.get("tiempo") or {}
    columna_fecha = _texto(tiempo_cfg.get("columna")) or "Fecha"
    horas = {HORA_DEL_DIA} | {str(h["hora"]) for h in tiempo_cfg.get("horas_aceptadas") or []}
    repetidas_cfg = (config.get("fechas_repetidas") or {}).get(nombre_hoja) or {}
    fuente = (config.get("indice") or {}).get("fuente")
    prov_id, geo_id, geo_nombre, geo_verificar = geo

    fijos = {
        "base_id": int(config["base"]), "entrega": entrega, "hoja": nombre_hoja,
        "ambito": hoja_cfg.get("ambito"), "grano_tiempo": "fecha",
        "provincia": (config.get("geo") or {}).get("particion", "sde"),
        "provincia_id": prov_id,
        "estacion_id": int(hoja_cfg["estacion_id"]), "estacion": hoja_cfg.get("estacion"),
        "estacion_nombre": estacion["estacion_nombre"],
        "estacion_tipo": estacion["estacion_tipo"],
        "estacion_id_interno": estacion["estacion_id_interno"],
        "localidad": estacion["localidad"], "ubicacion": estacion["ubicacion"],
        "latitud": estacion["latitud"], "longitud": estacion["longitud"],
        "altura": estacion["altura"],
        "estacion_desde_declarado": estacion["desde"],
        "estacion_hasta_declarado": estacion["hasta"],
        "nivel_geo": (config.get("geo") or {}).get("nivel", "estacion"),
        "geo_id": geo_id, "geo_nombre": geo_nombre, "geo_verificar": geo_verificar,
        "fuente": fuente,
    }

    pos_fecha, cols_var, vacias = None, None, None
    vistas = {}  # fecha -> {variable: (valor, valor_texto)} de la fila conservada
    for nro_fila, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
        if nro_fila < header_fila:
            continue
        if nro_fila == header_fila:
            pos_fecha, cols_var, vacias = _mapear_header(celdas, columna_fecha, variables,
                                                         nombre_hoja)
            continue
        celda_fecha = _celda(celdas, pos_fecha)
        if _es_vacia(celda_fecha):
            if all(_es_vacia(c) for c in celdas):
                continue  # fila en blanco al pie
            raise ValueError("Fila %d de la hoja %r: hay datos pero no hay fecha."
                             % (nro_fila, nombre_hoja))
        for pos in vacias:
            if not _es_vacia(_celda(celdas, pos)):
                raise ValueError(
                    "Fila %d de la hoja %r: la columna %d no tiene encabezado y trae %r."
                    % (nro_fila, nombre_hoja, pos + 1, _celda(celdas, pos))
                )
        fecha = _fecha_iso(celda_fecha, horas, nro_fila, nombre_hoja)
        valores = {}
        for pos, origen, canonico, meta in cols_var:
            valores[canonico] = _valor(_celda(celdas, pos), meta, nro_fila, origen,
                                       nombre_hoja)

        if fecha in vistas:
            _verificar_repetida(fecha, vistas[fecha], valores, repetidas_cfg, nro_fila,
                                nombre_hoja)
            continue  # declarada y sin contradiccion: se conserva la primera
        vistas[fecha] = valores

        for pos, origen, canonico, meta in cols_var:
            valor, valor_texto = valores[canonico]
            for clave, dato in fijos.items():
                columnas[clave].append(dato)
            columnas["fila_origen"].append(nro_fila)
            columnas["fecha"].append(fecha)
            columnas["anio"].append(fecha.year)
            columnas["mes"].append(fecha.month)
            columnas["dia"].append(fecha.day)
            columnas["variable"].append(canonico)
            columnas["variable_origen"].append(origen)
            columnas["variable_etiqueta"].append(meta.get("etiqueta"))
            columnas["unidad"].append(meta.get("unidad"))
            columnas["tipo_valor"].append(meta["tipo"])
            columnas["valor"].append(valor)
            columnas["valor_texto"].append(valor_texto)


def _mapear_header(celdas, columna_fecha, variables, nombre_hoja):
    """(pos_fecha, [(pos, origen, canonico, meta)], [posiciones sin encabezado])."""
    pos_fecha, cols_var, vacias, sin_mapeo, usados = None, [], [], [], {}
    for i, celda in enumerate(celdas):
        nombre = _texto(celda)
        if nombre is None:
            vacias.append(i)
            continue
        if nombre == columna_fecha:
            pos_fecha = i
            continue
        if nombre not in variables:
            sin_mapeo.append(nombre)
            continue
        canonico, meta = variables[nombre]
        if canonico in usados:
            raise ValueError(
                "Hoja %r: las columnas %r y %r van a la misma variable %r."
                % (nombre_hoja, usados[canonico], nombre, canonico)
            )
        usados[canonico] = nombre
        cols_var.append((i, nombre, canonico, meta))
    if sin_mapeo:
        raise ValueError(
            "Hoja %r: columnas que no estan en el diccionario de variables: %s. Agregalas a "
            "`columnas_origen` de la variable que corresponda (o una variable nueva) antes de "
            "ingerir." % (nombre_hoja, sin_mapeo)
        )
    if pos_fecha is None:
        raise ValueError("Hoja %r: no esta la columna de fecha %r." % (nombre_hoja,
                                                                       columna_fecha))
    return pos_fecha, cols_var, vacias


def _verificar_repetida(fecha, conservada, descartada, repetidas_cfg, nro_fila, nombre_hoja):
    decl = repetidas_cfg.get(fecha.isoformat())
    if not decl:
        raise ValueError(
            "Fila %d de la hoja %r: el dia %s ya aparecio antes. Declaralo en "
            "`fechas_repetidas` del config (con su motivo) o pedile a JC que lo corrija."
            % (nro_fila, nombre_hoja, fecha.isoformat())
        )
    if decl.get("conservar") != "primera":
        raise ValueError("`fechas_repetidas` de la hoja %r, dia %s: el unico criterio "
                         "soportado es `conservar: primera`." % (nombre_hoja, fecha))
    for variable, par in descartada.items():
        if par == (None, None):
            continue
        if conservada.get(variable) != par:
            raise ValueError(
                "Fila %d de la hoja %r (dia %s repetido): la variable %r dice %r y la primera "
                "fila del dia dice %r. Las dos filas se contradicen: no se descarta ninguna "
                "sin que JC decida." % (nro_fila, nombre_hoja, fecha, variable, par,
                                        conservada.get(variable))
            )


def _valor(celda, meta, nro_fila, origen, nombre_hoja):
    """(valor numerico, valor_texto) de una celda segun el tipo de la variable."""
    if _es_vacia(celda):
        return None, None
    if meta["tipo"] == "categoria":
        texto = str(celda).strip()
        if texto not in meta["dominio"]:
            raise ValueError(
                "Fila %d de la hoja %r: la columna %r trae %r, que no esta en el `dominio` de "
                "la variable." % (nro_fila, nombre_hoja, origen, celda)
            )
        return None, texto
    return _numero(celda, nro_fila, origen, nombre_hoja), None


# ---------------------------------------------------------------------------
# Diccionario de variables
# ---------------------------------------------------------------------------
def _cargar_variables(config):
    """{encabezado_origen: (canonico, meta)} del diccionario que declara el config."""
    ruta = config.get("variables")
    if not ruta:
        raise ValueError("Base %s: el config tiene que declarar `variables:` (la ruta del "
                         "diccionario de variables)." % config["base"])
    with open(os.path.join(RAIZ, ruta), encoding="utf-8") as f:
        diccionario = yaml.safe_load(f)
    por_origen = {}
    for canonico, meta in (diccionario.get("variables") or {}).items():
        meta = dict(meta or {})
        meta.setdefault("tipo", "numero")
        if meta["tipo"] not in TIPOS_DE_VALOR:
            raise ValueError("Variable %r: `tipo` tiene que ser uno de %s."
                             % (canonico, TIPOS_DE_VALOR))
        if meta["tipo"] == "categoria" and not meta.get("dominio"):
            raise ValueError("Variable %r: una variable de categoria declara su `dominio`."
                             % canonico)
        for origen in meta.get("columnas_origen") or []:
            if origen in por_origen:
                raise ValueError("El encabezado %r esta en dos variables: %r y %r."
                                 % (origen, por_origen[origen][0], canonico))
            por_origen[origen] = (canonico, meta)
    return por_origen


# ---------------------------------------------------------------------------
# Utilidades chicas
# ---------------------------------------------------------------------------
def _cargar_alias():
    with open(RUTA_ALIAS, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _normalizar(texto):
    if texto is None:
        return None
    sin_tildes = "".join(
        c for c in unicodedata.normalize("NFD", str(texto))
        if unicodedata.category(c) != "Mn"
    )
    return " ".join(sin_tildes.upper().split())


def _texto(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _es_vacia(celda):
    return celda is None or (isinstance(celda, str) and not celda.strip())


def _celda(celdas, pos):
    if pos is None:
        return None
    return celdas[pos] if pos < len(celdas) else None


def _entero_estricto(valor, nro_fila, campo, nombre_hoja):
    numero = _numero(valor, nro_fila, campo, nombre_hoja)
    if numero is None or numero != int(numero):
        raise ValueError("Fila %d de la hoja %r: %r no es un entero (%r)."
                         % (nro_fila, nombre_hoja, campo, valor))
    return int(numero)


def _numero(valor, nro_fila, campo, nombre_hoja):
    """float de una celda numerica o de un TEXTO con punto decimal ingles ('1.0E-7')."""
    if _es_vacia(valor):
        return None
    if isinstance(valor, bool):
        raise ValueError("Fila %d de la hoja %r: la columna %r trae un booleano."
                         % (nro_fila, nombre_hoja, campo))
    try:
        return float(str(valor).strip()) if isinstance(valor, str) else float(valor)
    except (TypeError, ValueError):
        raise ValueError(
            "Fila %d de la hoja %r: la columna %r trae %r, que no es un numero. No se "
            "convierte a cero: el dato se revisa en el Excel." % (nro_fila, nombre_hoja,
                                                                 campo, valor)
        )


def _fecha_iso(valor, horas, nro_fila, nombre_hoja):
    """Dia de una celda de fecha. `horas` son las horas del dia aceptadas ("00:00:00", ...)."""
    if isinstance(valor, dt.datetime):
        hora, dia = valor.strftime("%H:%M:%S"), valor.date()
    elif isinstance(valor, dt.date):
        return valor
    else:
        m = RE_FECHA.match(str(valor).strip())
        if not m:
            raise ValueError("Fila %d de la hoja %r: no se entiende la fecha %r."
                             % (nro_fila, nombre_hoja, valor))
        hora = m.group(4) or HORA_DEL_DIA
        dia = dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    if hora not in horas:
        raise ValueError(
            "Fila %d de la hoja %r: la fecha %r trae una hora que no es la del dia. Si es un "
            "artefacto de la exportacion, declarala en `tiempo.horas_aceptadas` del config con "
            "su motivo." % (nro_fila, nombre_hoja, valor)
        )
    return dia


def _fecha_dma(valor, nro_fila, nombre_hoja):
    """Fecha 'dd/mm/aaaa' de la hoja de estaciones (o nula si la celda esta vacia)."""
    if _es_vacia(valor):
        return None
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    m = RE_FECHA_DMA.match(str(valor).strip())
    if not m:
        raise ValueError("Fila %d de la hoja %r: no se entiende la fecha %r (se espera "
                         "dd/mm/aaaa)." % (nro_fila, nombre_hoja, valor))
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
