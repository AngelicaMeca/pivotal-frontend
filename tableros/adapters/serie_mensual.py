"""Adapter de la familia `serie-mensual` (lecheria: una fila por mes).

Dos formas de la misma familia, que decide el config con `apertura`:

  * `apertura: provincias` - la fila es un mes y hay UNA COLUMNA POR PROVINCIA, todas con la
    misma medida (el precio del litro, el porcentaje de grasa butirosa). Las bases 185, 186,
    188 y 194.
  * `apertura: medidas` - la fila es un mes de UNA sola provincia y cada columna es una medida
    distinta (litros, tambos, $/litro, %grasa). La base 184.

Por que no entra en las que ya estan:
  * `wide-mes` abre los MESES a lo ancho y aca el mes esta en las filas;
  * `stock` no tiene mes: su grano es el anio;
  * `precios` tiene la maquinaria de medidas no aditivas, que aca hace falta, pero su fila es
    una cotizacion suelta con el producto descrito en columnas, no una serie.

LAS MEDIDAS NO SE SUMAN. Un precio, un porcentaje de grasa y un promedio de litros por tambo
son no aditivos: cada medida declara su `agregacion` en el config y el adapter FALLA si no la
declara. La regla viaja pegada al dato (columna `agregacion`) para que ningun mart ni ninguna
vista tenga que adivinarla, igual que en la familia `precios`.

OJO CON LO QUE SE PUBLICA: la apertura por provincias es material de comparacion
interprovincial. Entra entero al mart -el dato es el dato- y que columnas se publican lo
decide el spec de cada vista, nunca el adapter.
"""
import datetime as dt
import os

import openpyxl
import polars as pl
import yaml

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUTA_ALIAS = os.path.join(RAIZ, "configs", "dims", "geo-alias.yaml")

AGREGACIONES = ("suma", "promedio", "ultimo")

COLUMNAS_SALIDA = [
    "base_id", "entrega", "hoja", "ambito", "fila_origen",
    "grano_tiempo", "anio", "mes",
    "provincia", "provincia_nombre", "provincia_id", "nivel_geo", "geo_id",
    "es_agregado_geo",
    "variable", "medida", "medida_etiqueta", "unidad", "agregable", "agregacion",
    "ponderacion", "valor", "fuente",
]

ESQUEMA_SALIDA = {
    "base_id": pl.Int64, "entrega": pl.Utf8, "hoja": pl.Utf8, "ambito": pl.Utf8,
    "fila_origen": pl.Int64, "grano_tiempo": pl.Utf8, "anio": pl.Int64, "mes": pl.Int64,
    "provincia": pl.Utf8, "provincia_nombre": pl.Utf8, "provincia_id": pl.Int64,
    "nivel_geo": pl.Utf8, "geo_id": pl.Utf8, "es_agregado_geo": pl.Boolean,
    "variable": pl.Utf8, "medida": pl.Utf8, "medida_etiqueta": pl.Utf8, "unidad": pl.Utf8,
    "agregable": pl.Boolean, "agregacion": pl.Utf8, "ponderacion": pl.Utf8,
    "valor": pl.Float64, "fuente": pl.Utf8,
}

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}


def parse(config, path, entrega=None):
    alias = _cargar_alias()
    libro = openpyxl.load_workbook(path, read_only=True, data_only=True)
    filas = []
    try:
        for hoja_cfg in (config.get("hojas") or {}).get("datos") or []:
            filas.extend(_parsear_hoja(config, libro, hoja_cfg, entrega, alias))
    finally:
        libro.close()
    if not filas:
        return pl.DataFrame(schema=ESQUEMA_SALIDA)
    return pl.DataFrame(filas, schema=ESQUEMA_SALIDA).select(COLUMNAS_SALIDA)


def _parsear_hoja(config, libro, hoja_cfg, entrega, alias):
    nombre_hoja = hoja_cfg["nombre"]
    if nombre_hoja not in libro.sheetnames:
        raise ValueError(
            "La hoja %r que declara el config no esta en el archivo. Hojas presentes: %s"
            % (nombre_hoja, libro.sheetnames)
        )
    hoja = libro[nombre_hoja]
    header_fila = int(hoja_cfg.get("header_fila", 1))
    apertura = config.get("apertura")
    if apertura not in ("provincias", "medidas"):
        raise ValueError(
            "El config tiene que declarar `apertura: provincias` o `apertura: medidas`; "
            "trae %r." % apertura
        )

    mapeo = {_texto(k): v for k, v in (config.get("columnas") or {}).items()}
    mapeo.update({_texto(k): v for k, v in (hoja_cfg.get("columnas") or {}).items()})
    medida_cfg = hoja_cfg.get("medida") or config.get("medida") or {}
    valores = config.get("valores") or {}
    fuente = (config.get("indice") or {}).get("fuente")

    filas = []
    posiciones, provincias = None, None
    for nro_fila, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
        if nro_fila < header_fila:
            continue
        if nro_fila == header_fila:
            posiciones, provincias = _mapear_header(celdas, mapeo, nombre_hoja, apertura,
                                                    alias, config)
            continue
        if all(_es_vacia(c) for c in celdas):
            continue
        anio, mes = _periodo(_celda(celdas, posiciones.get("periodo")),
                             _celda(celdas, posiciones.get("mes")), nro_fila, nombre_hoja)
        if anio is None:
            continue  # pie de tabla ("Fuente: ..."), no es una observacion
        if apertura == "provincias":
            filas.extend(_filas_por_provincia(config, hoja_cfg, celdas, provincias, anio,
                                              mes, nro_fila, entrega, medida_cfg, fuente))
        else:
            filas.extend(_filas_por_medida(config, hoja_cfg, celdas, posiciones, valores,
                                           anio, mes, nro_fila, entrega, fuente, alias))
    return filas


def _mapear_header(celdas, mapeo, nombre_hoja, apertura, alias, config):
    """{campo: posicion} y, con `apertura: provincias`, {posicion: (nombre, id)}."""
    posiciones, provincias = {}, {}
    sin_mapeo = []
    por_nombre = alias.get("provincia") or {}
    indice = {_normalizar(k): (k, v) for k, v in por_nombre.items()}
    for i, celda in enumerate(celdas):
        nombre = _texto(celda)
        if nombre is None:
            continue
        if nombre in mapeo:
            destino = mapeo[nombre]
            if destino is not None:
                posiciones[destino] = i
            continue
        if apertura == "provincias":
            # Toda columna que no este mapeada tiene que ser una PROVINCIA. Si no se
            # reconoce, corta: una columna nueva en el origen no se pierde en silencio.
            clave = _normalizar(nombre)
            if clave in indice:
                provincias[i] = indice[clave]
                continue
            if clave in ("NACIONAL", "TOTAL PAIS", "PAIS"):
                provincias[i] = ("Total del pais", None)
                continue
        sin_mapeo.append(nombre)
    if sin_mapeo:
        raise ValueError(
            "La hoja %r tiene columnas sin mapeo en el config: %s. Agregalas a `columnas` "
            "(o a geo-alias.yaml si son provincias) antes de ingerir."
            % (nombre_hoja, sin_mapeo)
        )
    if "periodo" not in posiciones:
        raise ValueError(
            "La hoja %r no tiene la columna de periodo. Mapeala a `periodo` en `columnas`."
            % nombre_hoja
        )
    return posiciones, provincias


def _filas_por_provincia(config, hoja_cfg, celdas, provincias, anio, mes, nro_fila,
                         entrega, medida_cfg, fuente):
    _validar_agregacion(medida_cfg, config["base"])
    quitar = tuple((config.get("limpieza") or {}).get("quitar_prefijos") or ())
    salida = []
    for pos, (nombre_prov, prov_id) in sorted(provincias.items()):
        valor = _numero(_celda(celdas, pos), nro_fila, nombre_prov, hoja_cfg["nombre"],
                        quitar)
        salida.append({
            "base_id": int(config["base"]), "entrega": entrega,
            "hoja": hoja_cfg["nombre"], "ambito": hoja_cfg.get("ambito"),
            "fila_origen": nro_fila, "grano_tiempo": "anio_mes",
            "anio": anio, "mes": mes,
            "provincia": (config.get("geo") or {}).get("particion", "sde"),
            "provincia_nombre": nombre_prov,
            "provincia_id": prov_id,
            "nivel_geo": "provincia",
            # Sin geo_id: la familia no aporta departamentos a dim_geo.
            "geo_id": None,
            "es_agregado_geo": True,
            "variable": medida_cfg.get("variable", "valor"),
            "medida": medida_cfg.get("medida"),
            "medida_etiqueta": medida_cfg.get("etiqueta"),
            "unidad": medida_cfg.get("unidad"),
            "agregable": bool(medida_cfg.get("agregable", False)),
            "agregacion": medida_cfg.get("agregacion"),
            "ponderacion": medida_cfg.get("ponderacion"),
            "valor": valor, "fuente": fuente,
        })
    return salida


def _filas_por_medida(config, hoja_cfg, celdas, posiciones, valores, anio, mes, nro_fila,
                      entrega, fuente, alias):
    geo_cfg = config.get("geo") or {}
    nombre_prov = _texto(geo_cfg.get("provincia_fija"))
    prov_id = (alias.get("provincia") or {}).get(nombre_prov)
    if prov_id is None:
        raise ValueError(
            "La provincia %r de `geo.provincia_fija` no esta en geo-alias.yaml." % nombre_prov
        )
    quitar = tuple((config.get("limpieza") or {}).get("quitar_prefijos") or ())
    salida = []
    for campo, meta in valores.items():
        if campo not in posiciones:
            continue
        _validar_agregacion(meta, config["base"], campo)
        salida.append({
            "base_id": int(config["base"]), "entrega": entrega,
            "hoja": hoja_cfg["nombre"], "ambito": hoja_cfg.get("ambito"),
            "fila_origen": nro_fila, "grano_tiempo": "anio_mes",
            "anio": anio, "mes": mes,
            "provincia": geo_cfg.get("particion", "sde"),
            "provincia_nombre": nombre_prov, "provincia_id": prov_id,
            "nivel_geo": "provincia", "geo_id": None, "es_agregado_geo": True,
            "variable": campo,
            "medida": meta.get("medida"), "medida_etiqueta": meta.get("etiqueta"),
            "unidad": meta.get("unidad"),
            "agregable": bool(meta.get("agregable", False)),
            "agregacion": meta.get("agregacion"),
            "ponderacion": meta.get("ponderacion"),
            "valor": _numero(_celda(celdas, posiciones[campo]), nro_fila, campo,
                             hoja_cfg["nombre"], quitar),
            "fuente": fuente,
        })
    return salida


def _validar_agregacion(meta, base, campo="la medida"):
    """Ninguna medida de esta familia se agrega sin decir COMO. Es la regla de `precios`."""
    agregacion = meta.get("agregacion")
    if agregacion not in AGREGACIONES:
        raise ValueError(
            "Base %s: %s no declara `agregacion` (%s). Las medidas de esta familia no son "
            "aditivas y el mart no puede adivinar como se agregan."
            % (base, campo, " | ".join(AGREGACIONES))
        )


def _periodo(celda_periodo, celda_mes, nro_fila, nombre_hoja):
    """(anio, mes) de la columna de periodo. Acepta fecha, 'AAAA-MM' o anio + nombre de mes."""
    if isinstance(celda_periodo, (dt.datetime, dt.date)):
        return int(celda_periodo.year), int(celda_periodo.month)
    anio = _entero(celda_periodo)
    if anio is not None and 1900 < anio < 2200:
        mes = _entero(celda_mes)
        if mes is None:
            mes = MESES.get(_normalizar(celda_mes).lower() if celda_mes else "")
        if mes is None:
            raise ValueError(
                "Fila %d de la hoja %r: el anio es %s pero no se entiende el mes %r."
                % (nro_fila, nombre_hoja, anio, celda_mes)
            )
        return anio, mes
    texto = _texto(celda_periodo)
    if texto and len(texto) >= 7 and texto[:4].isdigit() and texto[5:7].isdigit():
        return int(texto[:4]), int(texto[5:7])
    return None, None


# ---------------------------------------------------------------------------
# Utilidades chicas
# ---------------------------------------------------------------------------
def _cargar_alias():
    with open(RUTA_ALIAS, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _normalizar(texto):
    import unicodedata
    if texto is None:
        return None
    sin_tildes = "".join(
        c for c in unicodedata.normalize("NFD", str(texto))
        if unicodedata.category(c) != "Mn"
    )
    return " ".join(sin_tildes.upper().split())


def _texto(valor):
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _es_vacia(celda):
    return celda is None or (isinstance(celda, str) and not celda.strip())


def _celda(celdas, pos):
    if pos is None:
        return None
    return celdas[pos] if pos < len(celdas) else None


def _entero(valor):
    if _es_vacia(valor):
        return None
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return None


def _numero(valor, nro_fila, campo, nombre_hoja, quitar=()):
    if _es_vacia(valor):
        return None
    if isinstance(valor, bool):
        raise ValueError(
            "Fila %d de la hoja %r: la columna %r trae un booleano."
            % (nro_fila, nombre_hoja, campo)
        )
    # Celdas que la planilla guardo como TEXTO con el simbolo de moneda adelante ("$ 235.39").
    # No es un dato roto, es un formato: el numero se lee sin ambiguedad. Los prefijos que se
    # sacan los DECLARA el config uno por uno; cualquier otro texto sigue cortando la ingesta.
    if isinstance(valor, str) and quitar:
        limpio = valor.strip()
        for prefijo in quitar:
            if limpio.startswith(prefijo):
                limpio = limpio[len(prefijo):].strip()
        valor = limpio
    try:
        return float(valor)
    except (TypeError, ValueError):
        raise ValueError(
            "Fila %d de la hoja %r: la columna %r trae %r, que no es un numero. "
            "No se convierte a cero: el dato se revisa en el Excel."
            % (nro_fila, nombre_hoja, campo, valor)
        )
